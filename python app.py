import hashlib
import json
import os
import re
import socket
import sqlite3
import time as pytime
import traceback
from pathlib import Path
from datetime import date, datetime, time, timedelta, timezone
from functools import wraps
from typing import Any
from urllib.parse import urlsplit, urlunsplit
from uuid import uuid4

from authlib.integrations.flask_client import OAuth
from flask import Flask, render_template, request, jsonify, session, redirect, url_for, Response, abort
from openai import OpenAI
from werkzeug.utils import secure_filename

try:
    from flask_socketio import SocketIO, emit, join_room
except Exception:  # pragma: no cover
    SocketIO = None  # type: ignore[assignment]
    emit = None  # type: ignore[assignment]
    join_room = None  # type: ignore[assignment]

try:
    from zoneinfo import ZoneInfo  # py3.9+
except Exception:  # pragma: no cover
    ZoneInfo = None

try:
    import stripe  # type: ignore
except Exception:  # pragma: no cover
    stripe = None  # type: ignore

try:
    import firebase_admin  # type: ignore
    from firebase_admin import credentials as fb_credentials  # type: ignore
    from firebase_admin import auth as fb_auth  # type: ignore
    from firebase_admin import firestore as fb_firestore  # type: ignore
except Exception:  # pragma: no cover
    firebase_admin = None  # type: ignore
    fb_credentials = None  # type: ignore
    fb_auth = None  # type: ignore
    fb_firestore = None  # type: ignore

try:
    # firebase-admin bundles google-cloud-firestore; we import it lazily for transactions.
    from google.cloud import firestore as gc_firestore  # type: ignore
except Exception:  # pragma: no cover
    gc_firestore = None  # type: ignore

app = Flask(__name__)
app.secret_key = os.getenv("FLASK_SECRET_KEY", "dev-insecure-secret-change-me")
app.config["SESSION_COOKIE_SAMESITE"] = "Lax"
app.config["SESSION_COOKIE_HTTPONLY"] = True
# If you deploy behind HTTPS, set this env var to "1" so cookies are secure.
# In local dev (HTTP on LAN IP), Secure cookies would break login/session on phones.
_debug_env = (os.getenv("FLASK_DEBUG", "0") == "1")
app.config["SESSION_COOKIE_SECURE"] = (os.getenv("COOKIE_SECURE", "0") == "1") and (not _debug_env)
app.config["TEMPLATES_AUTO_RELOAD"] = True
# Avoid stale CSS/JS during local development.
app.config["SEND_FILE_MAX_AGE_DEFAULT"] = 0

BASE_DIR = os.path.dirname(os.path.abspath(__file__))


def _load_dotenv_if_present():
    """
    Lightweight .env loader so running via PyCharm (or `python "python app.py"`) still works.
    Does not override already-set environment variables.
    """
    path = os.path.join(BASE_DIR, ".env")
    if not os.path.exists(path):
        return
    try:
        with open(path, "r", encoding="utf-8") as f:
            for raw in f:
                line = (raw or "").strip()
                if not line or line.startswith("#"):
                    continue
                if "=" not in line:
                    continue
                k, v = line.split("=", 1)
                k = (k or "").strip()
                v = (v or "").strip()
                if not k or k in os.environ:
                    continue
                if len(v) >= 2 and v.startswith('"') and v.endswith('"'):
                    v = v[1:-1]
                os.environ[k] = v
    except Exception:
        return


_load_dotenv_if_present()

@app.context_processor 
def _inject_static_v(): 
    """ 
    Cache-bust CSS/JS during local UI iterations so browsers don't keep stale assets. 
 
    Behavior: 
    - Default: always use a per-request timestamp (development-friendly, avoids stale CSS/JS). 
    - Optional: set STATIC_V_MODE=env (or "forced") to use STATIC_V as a stable version string. 
    """ 
    mode = str(os.getenv("STATIC_V_MODE", "") or "").strip().lower() 
    if mode in ("env", "forced"): 
        forced = str(os.getenv("STATIC_V", "") or "").strip() 
        if forced: 
            return {"static_v": forced} 
 
    return {"static_v": str(int(pytime.time()))} 

@app.context_processor
def _inject_csrf():
    # Minimal CSRF token for HTML form posts.
    tok = session.get("csrf_token")
    if not tok:
        tok = uuid4().hex
        session["csrf_token"] = tok
    return {"csrf_token": tok}

@app.context_processor
def _inject_user_ctx():
    """
    Common user-facing identity values for templates.
    - `username` is the local stable identifier.
    - `display_name` is the friendly nickname shown in UI.
    """
    username = str(session.get("username", "") or "")
    display_name = str(session.get("display_name", "") or "")
    email = str(session.get("email", "") or "")
    if not display_name:
        # Older sessions (or legacy logins) may not have display_name yet.
        if ("@" in email) and (not re.search(r"\s", email)):
            display_name = email.split("@", 1)[0]
        else:
            display_name = username

    return {
        "username": username,
        "display_name": display_name,
        "is_guest": bool(session.get("is_guest")),
        "sign_in_provider": str(session.get("sign_in_provider", "") or ""),
    }


DB_PATH = os.path.join(BASE_DIR, "users.db")
conversations = {}
UPLOAD_FOLDER = os.path.join(BASE_DIR, "static", "uploads")
ALLOWED_IMAGE_EXTENSIONS = {"png", "jpg", "jpeg", "gif", "webp"}
CHAT_MODES = {
    "normal": "You are yan, a helpful assistant.",
    # Next-gen modes
    "creative": "You are yan in Creative Mode. Be imaginative, playful, and practical. Offer multiple options and bold ideas.",
    "coding": "You are yan in Coding Mode. Be precise, include correct code, explain tradeoffs, and suggest tests and edge cases.",
    "research": "You are yan in Research Mode. Ask clarifying questions, be careful with uncertainty, and structure the answer with claims and caveats. If you cannot browse, say so.",
    "learning": "You are yan in Learning Mode. Teach step-by-step, check understanding, and include short exercises or examples.",
    "business": "You are yan in Business Mode. Be concise, decision-oriented, and focus on ROI, risks, and next steps.",
    "translator": "You are yan in Translator Mode. Translate and improve text naturally, keeping meaning intact.",
    # Back-compat aliases
    "teacher": "You are yan in Learning Mode. Teach step-by-step, check understanding, and include short exercises or examples.",
    "brainstorm": "You are yan in Creative Mode. Be imaginative, playful, and practical. Offer multiple options and bold ideas.",
}

oauth = OAuth(app)
oauth_clients = {}

socketio = None
if SocketIO is not None:
    # Use threading mode by default so this works without eventlet/gevent.
    socketio = SocketIO(app, async_mode="threading", cors_allowed_origins=None)

FREE_PLAN = "free"
PRO_PLANS = {"pro", "ultimate", "paid", "premium"}


def _parse_int_env(name: str, default: int) -> int:
    raw = (os.getenv(name) or "").strip()
    if not raw:
        return default
    try:
        return int(raw)
    except ValueError:
        return default


def client_ip() -> str:
    # Only trust X-Forwarded-For if you're behind a trusted reverse proxy.
    ip = (request.headers.get("X-Forwarded-For") or "").split(",")[0].strip()
    if not ip:
        ip = (request.remote_addr or "").strip()
    return ip or "unknown"


def require_csrf():
    # Enforce only for HTML form submissions to avoid breaking JSON APIs.
    if request.method != "POST":
        return
    ct = (request.content_type or "").lower()
    if ("application/x-www-form-urlencoded" not in ct) and ("multipart/form-data" not in ct):
        return
    sent = (request.form.get("csrf_token") or "").strip()
    tok = (session.get("csrf_token") or "").strip()
    if not tok or not sent or sent != tok:
        abort(400, description="Invalid CSRF token.")


# Persist login across browser restarts by default.
app.config["PERMANENT_SESSION_LIFETIME"] = timedelta(days=max(1, _parse_int_env("SESSION_DAYS", 30)))


FREE_DAILY_MESSAGE_LIMIT = max(0, _parse_int_env("FREE_DAILY_MESSAGE_LIMIT", 15))


def get_openai_client():
    key = (os.getenv("OPENAI_API_KEY") or "").strip()
    if not key:
        return None
    return OpenAI(api_key=key)


def init_db():
    conn = sqlite3.connect(DB_PATH)
    cur = conn.cursor()
    cur.execute(
        """
        CREATE TABLE IF NOT EXISTS users (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            username TEXT UNIQUE NOT NULL,
            password_hash TEXT NOT NULL
        )
        """
    )
    # Case-insensitive uniqueness for usernames.
    # If the DB already contains usernames that differ only by case, this will fail; we ignore it.
    try:
        cur.execute("CREATE UNIQUE INDEX IF NOT EXISTS idx_users_username_nocase ON users(username COLLATE NOCASE)")
    except Exception:
        pass
    cur.execute(
        """
        CREATE TABLE IF NOT EXISTS identities (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            user_id INTEGER NOT NULL,
            provider TEXT NOT NULL,
            provider_user_id TEXT NOT NULL,
            email TEXT,
            UNIQUE(provider, provider_user_id),
            FOREIGN KEY(user_id) REFERENCES users(id)
        )
        """
    )
    cur.execute("PRAGMA table_info(users)")
    user_columns = {row[1] for row in cur.fetchall()}
    if "email" not in user_columns:
        cur.execute("ALTER TABLE users ADD COLUMN email TEXT")
    if "address" not in user_columns:
        cur.execute("ALTER TABLE users ADD COLUMN address TEXT")
    if "age" not in user_columns:
        cur.execute("ALTER TABLE users ADD COLUMN age INTEGER")
    if "photo_path" not in user_columns:
        cur.execute("ALTER TABLE users ADD COLUMN photo_path TEXT")
    if "plan" not in user_columns:
        cur.execute("ALTER TABLE users ADD COLUMN plan TEXT")
    if "tz" not in user_columns:
        cur.execute("ALTER TABLE users ADD COLUMN tz TEXT")
    if "stripe_customer_id" not in user_columns:
        cur.execute("ALTER TABLE users ADD COLUMN stripe_customer_id TEXT")
    if "created_at_utc" not in user_columns:
        cur.execute("ALTER TABLE users ADD COLUMN created_at_utc TEXT")
    if "last_login_at_utc" not in user_columns:
        cur.execute("ALTER TABLE users ADD COLUMN last_login_at_utc TEXT")
    if "last_login_ip" not in user_columns:
        cur.execute("ALTER TABLE users ADD COLUMN last_login_ip TEXT")
    if "failed_login_count" not in user_columns:
        cur.execute("ALTER TABLE users ADD COLUMN failed_login_count INTEGER")
    if "locked_until_utc" not in user_columns:
        cur.execute("ALTER TABLE users ADD COLUMN locked_until_utc TEXT")
    if "password_changed_at_utc" not in user_columns:
        cur.execute("ALTER TABLE users ADD COLUMN password_changed_at_utc TEXT")
    if "daily_limit_override" not in user_columns:
        cur.execute("ALTER TABLE users ADD COLUMN daily_limit_override INTEGER")
    cur.execute("UPDATE users SET plan = ? WHERE plan IS NULL OR TRIM(plan) = ''", (FREE_PLAN,))
    now = utc_now_iso()
    cur.execute("UPDATE users SET created_at_utc = ? WHERE created_at_utc IS NULL OR TRIM(created_at_utc) = ''", (now,))
    cur.execute("UPDATE users SET failed_login_count = 0 WHERE failed_login_count IS NULL")
    cur.execute("UPDATE users SET locked_until_utc = '' WHERE locked_until_utc IS NULL")
    cur.execute("UPDATE users SET daily_limit_override = NULL WHERE daily_limit_override IS NOT NULL AND daily_limit_override < 0")
    cur.execute(
        "UPDATE users SET password_changed_at_utc = COALESCE(created_at_utc, ?) "
        "WHERE password_changed_at_utc IS NULL OR TRIM(password_changed_at_utc) = ''",
        (now,),
    )
    cur.execute("CREATE INDEX IF NOT EXISTS idx_users_email_nocase ON users(email COLLATE NOCASE)")
    cur.execute("CREATE INDEX IF NOT EXISTS idx_users_last_login ON users(last_login_at_utc)")

    cur.execute(
        """
        CREATE TABLE IF NOT EXISTS auth_attempts (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            action TEXT NOT NULL,
            key TEXT NOT NULL,
            ip TEXT NOT NULL,
            created_at_utc TEXT NOT NULL
        )
        """
    )
    cur.execute("CREATE INDEX IF NOT EXISTS idx_auth_attempts_action_ip_time ON auth_attempts(action, ip, created_at_utc)")
    cur.execute("CREATE INDEX IF NOT EXISTS idx_auth_attempts_action_key_time ON auth_attempts(action, key, created_at_utc)")

    cur.execute(
        """
        CREATE TABLE IF NOT EXISTS password_reset_tokens (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            user_id INTEGER NOT NULL,
            token_hash TEXT NOT NULL,
            created_at_utc TEXT NOT NULL,
            expires_at_utc TEXT NOT NULL,
            used_at_utc TEXT NOT NULL DEFAULT '',
            requested_ip TEXT NOT NULL DEFAULT '',
            FOREIGN KEY(user_id) REFERENCES users(id)
        )
        """
    )
    cur.execute("CREATE INDEX IF NOT EXISTS idx_pwreset_token_hash ON password_reset_tokens(token_hash)")
    cur.execute("CREATE INDEX IF NOT EXISTS idx_pwreset_user ON password_reset_tokens(user_id, created_at_utc)")
    cur.execute(
        """
        CREATE TABLE IF NOT EXISTS daily_message_usage (
            user_id INTEGER NOT NULL,
            day TEXT NOT NULL,
            message_count INTEGER NOT NULL DEFAULT 0,
            PRIMARY KEY (user_id, day),
            FOREIGN KEY(user_id) REFERENCES users(id)
        )
        """
    )
    cur.execute(
        """
        CREATE TABLE IF NOT EXISTS feedback (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            created_at_utc TEXT NOT NULL,
            user_id INTEGER,
            username TEXT,
            contact_email TEXT,
            rating INTEGER,
            category TEXT,
            message TEXT NOT NULL,
            page TEXT,
            tz TEXT,
            user_agent TEXT,
            ip TEXT,
            FOREIGN KEY(user_id) REFERENCES users(id)
        )
        """
    )

    # Instagram-like follow graph.
    cur.execute(
        """
        CREATE TABLE IF NOT EXISTS user_follows (
            follower_id INTEGER NOT NULL,
            followed_id INTEGER NOT NULL,
            created_at_utc TEXT NOT NULL,
            PRIMARY KEY (follower_id, followed_id),
            FOREIGN KEY(follower_id) REFERENCES users(id),
            FOREIGN KEY(followed_id) REFERENCES users(id)
        )
        """
    )
    cur.execute("CREATE INDEX IF NOT EXISTS idx_user_follows_follower ON user_follows(follower_id, followed_id)")
    cur.execute("CREATE INDEX IF NOT EXISTS idx_user_follows_followed ON user_follows(followed_id, follower_id)")

    # Messaging tables (1:1 and (optional) group chats).
    cur.execute(
        """
        CREATE TABLE IF NOT EXISTS chat_conversations (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            is_group INTEGER NOT NULL DEFAULT 0,
            title TEXT,
            created_at_utc TEXT NOT NULL
        )
        """
    )
    cur.execute(
        """
        CREATE TABLE IF NOT EXISTS chat_members (
            conversation_id INTEGER NOT NULL,
            user_id INTEGER NOT NULL,
            role TEXT,
            joined_at_utc TEXT NOT NULL,
            PRIMARY KEY (conversation_id, user_id),
            FOREIGN KEY(conversation_id) REFERENCES chat_conversations(id),
            FOREIGN KEY(user_id) REFERENCES users(id)
        )
        """
    )
    cur.execute(
        """
        CREATE TABLE IF NOT EXISTS chat_messages (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            conversation_id INTEGER NOT NULL,
            sender_id INTEGER NOT NULL,
            type TEXT NOT NULL DEFAULT 'text',
            body TEXT,
            media_path TEXT,
            created_at_utc TEXT NOT NULL,
            FOREIGN KEY(conversation_id) REFERENCES chat_conversations(id),
            FOREIGN KEY(sender_id) REFERENCES users(id)
        )
        """
    )
    cur.execute(
        """
        CREATE TABLE IF NOT EXISTS chat_reads (
            message_id INTEGER NOT NULL,
            user_id INTEGER NOT NULL,
            read_at_utc TEXT NOT NULL,
            PRIMARY KEY (message_id, user_id),
            FOREIGN KEY(message_id) REFERENCES chat_messages(id),
            FOREIGN KEY(user_id) REFERENCES users(id)
        )
        """
    )
    cur.execute("CREATE INDEX IF NOT EXISTS idx_chat_messages_conv_id ON chat_messages(conversation_id, id)")
    cur.execute("CREATE INDEX IF NOT EXISTS idx_chat_members_user_id ON chat_members(user_id, conversation_id)")

    # AI chat tables (single-user conversations with assistant messages).
    cur.execute(
        """
        CREATE TABLE IF NOT EXISTS ai_conversations (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            user_id INTEGER NOT NULL,
            title TEXT NOT NULL,
            mode TEXT NOT NULL DEFAULT 'normal',
            system_prompt TEXT NOT NULL DEFAULT '',
            is_pinned INTEGER NOT NULL DEFAULT 0,
            is_bookmarked INTEGER NOT NULL DEFAULT 0,
            summary TEXT NOT NULL DEFAULT '',
            created_at_utc TEXT NOT NULL,
            updated_at_utc TEXT NOT NULL,
            FOREIGN KEY(user_id) REFERENCES users(id)
        )
        """
    )
    cur.execute("CREATE INDEX IF NOT EXISTS idx_ai_conversations_user_updated ON ai_conversations(user_id, updated_at_utc)")
    cur.execute(
        """
        CREATE TABLE IF NOT EXISTS ai_messages (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            conversation_id INTEGER NOT NULL,
            role TEXT NOT NULL,
            content TEXT NOT NULL,
            created_at_utc TEXT NOT NULL,
            edited_at_utc TEXT,
            FOREIGN KEY(conversation_id) REFERENCES ai_conversations(id)
        )
        """
    )
    cur.execute("CREATE INDEX IF NOT EXISTS idx_ai_messages_conv_id ON ai_messages(conversation_id, id)")

    # If this instance already created ai_conversations without newer columns, add them.
    cur.execute("PRAGMA table_info(ai_conversations)")
    ai_conv_cols = {row[1] for row in cur.fetchall()}
    if "summary" not in ai_conv_cols:
        cur.execute("ALTER TABLE ai_conversations ADD COLUMN summary TEXT NOT NULL DEFAULT ''")

    # Workspace features: notes, memory, prompts, files, sharing.
    cur.execute(
        """
        CREATE TABLE IF NOT EXISTS ai_notes (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            conversation_id INTEGER NOT NULL,
            user_id INTEGER NOT NULL,
            body TEXT NOT NULL DEFAULT '',
            updated_at_utc TEXT NOT NULL,
            UNIQUE(conversation_id, user_id),
            FOREIGN KEY(conversation_id) REFERENCES ai_conversations(id),
            FOREIGN KEY(user_id) REFERENCES users(id)
        )
        """
    )
    cur.execute("CREATE INDEX IF NOT EXISTS idx_ai_notes_user_id ON ai_notes(user_id, conversation_id)")

    cur.execute(
        """
        CREATE TABLE IF NOT EXISTS user_memories (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            user_id INTEGER NOT NULL,
            text TEXT NOT NULL,
            created_at_utc TEXT NOT NULL,
            source_conversation_id INTEGER,
            FOREIGN KEY(user_id) REFERENCES users(id)
        )
        """
    )
    cur.execute("CREATE INDEX IF NOT EXISTS idx_user_memories_user_id ON user_memories(user_id, id)")

    cur.execute(
        """
        CREATE TABLE IF NOT EXISTS ai_prompts (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            user_id INTEGER NOT NULL,
            title TEXT NOT NULL,
            body TEXT NOT NULL,
            created_at_utc TEXT NOT NULL,
            updated_at_utc TEXT NOT NULL,
            FOREIGN KEY(user_id) REFERENCES users(id)
        )
        """
    )
    cur.execute("CREATE INDEX IF NOT EXISTS idx_ai_prompts_user_id ON ai_prompts(user_id, updated_at_utc)")

    cur.execute(
        """
        CREATE TABLE IF NOT EXISTS ai_files (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            conversation_id INTEGER NOT NULL,
            user_id INTEGER NOT NULL,
            filename TEXT NOT NULL,
            mime TEXT NOT NULL,
            size_bytes INTEGER NOT NULL,
            storage_path TEXT NOT NULL,
            created_at_utc TEXT NOT NULL,
            FOREIGN KEY(conversation_id) REFERENCES ai_conversations(id),
            FOREIGN KEY(user_id) REFERENCES users(id)
        )
        """
    )
    cur.execute("CREATE INDEX IF NOT EXISTS idx_ai_files_conv_id ON ai_files(conversation_id, id)")

    cur.execute(
        """
        CREATE TABLE IF NOT EXISTS ai_shares (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            conversation_id INTEGER NOT NULL,
            user_id INTEGER NOT NULL,
            token TEXT UNIQUE NOT NULL,
            created_at_utc TEXT NOT NULL,
            revoked_at_utc TEXT NOT NULL DEFAULT '',
            FOREIGN KEY(conversation_id) REFERENCES ai_conversations(id),
            FOREIGN KEY(user_id) REFERENCES users(id)
        )
        """
    )
    cur.execute("CREATE INDEX IF NOT EXISTS idx_ai_shares_user_id ON ai_shares(user_id, conversation_id)")

    # Stripe subscription state (one active subscription per user for now).
    cur.execute(
        """
        CREATE TABLE IF NOT EXISTS billing_subscriptions (
            user_id INTEGER PRIMARY KEY,
            provider TEXT NOT NULL DEFAULT 'stripe',
            customer_id TEXT NOT NULL DEFAULT '',
            subscription_id TEXT NOT NULL DEFAULT '',
            price_id TEXT NOT NULL DEFAULT '',
            status TEXT NOT NULL DEFAULT '',
            current_period_end_utc TEXT NOT NULL DEFAULT '',
            cancel_at_period_end INTEGER NOT NULL DEFAULT 0,
            plan TEXT NOT NULL DEFAULT '',
            interval TEXT NOT NULL DEFAULT '',
            updated_at_utc TEXT NOT NULL DEFAULT '',
            FOREIGN KEY(user_id) REFERENCES users(id)
        )
        """
    )
    cur.execute("CREATE INDEX IF NOT EXISTS idx_billing_subscriptions_customer ON billing_subscriptions(customer_id)")
    os.makedirs(UPLOAD_FOLDER, exist_ok=True)
    conn.commit()
    conn.close()


def _firebase_web_config() -> dict[str, str]:
    cfg = {
        "apiKey": (os.getenv("FIREBASE_API_KEY") or "").strip(),
        "authDomain": (os.getenv("FIREBASE_AUTH_DOMAIN") or "").strip(),
        "projectId": (os.getenv("FIREBASE_PROJECT_ID") or "").strip(),
        "appId": (os.getenv("FIREBASE_APP_ID") or "").strip(),
        "messagingSenderId": (os.getenv("FIREBASE_MESSAGING_SENDER_ID") or "").strip(),
        "storageBucket": (os.getenv("FIREBASE_STORAGE_BUCKET") or "").strip(),
    }
    return {k: v for k, v in cfg.items() if v}


def _firebase_init_error() -> str:
    if firebase_admin is None or fb_credentials is None or fb_auth is None:
        return "firebase-admin is not installed on the server."
    if firebase_admin._apps:  # type: ignore[attr-defined]
        return ""

    cred_path = (os.getenv("FIREBASE_ADMIN_CREDENTIALS_PATH") or "").strip()
    cred_json = (os.getenv("FIREBASE_ADMIN_CREDENTIALS_JSON") or "").strip()

    try:
        if cred_path:
            if not os.path.isabs(cred_path):
                cred_path = os.path.join(BASE_DIR, cred_path)
            cred = fb_credentials.Certificate(cred_path)
            firebase_admin.initialize_app(cred)
            return ""
        if cred_json:
            cred = fb_credentials.Certificate(json.loads(cred_json))
            firebase_admin.initialize_app(cred)
            return ""
        # Allow GOOGLE_APPLICATION_CREDENTIALS to be used if already configured.
        if (os.getenv("GOOGLE_APPLICATION_CREDENTIALS") or "").strip():
            firebase_admin.initialize_app()
            return ""
        return "Missing Firebase admin credentials. Set FIREBASE_ADMIN_CREDENTIALS_PATH or FIREBASE_ADMIN_CREDENTIALS_JSON."
    except Exception as exc:
        return f"Firebase init failed: {str(exc)}"


def configure_oauth():
    google_client_id = os.getenv("GOOGLE_CLIENT_ID")
    google_client_secret = os.getenv("GOOGLE_CLIENT_SECRET")
    ms_client_id = os.getenv("MICROSOFT_CLIENT_ID")
    ms_client_secret = os.getenv("MICROSOFT_CLIENT_SECRET")

    if google_client_id and google_client_secret:
        oauth_clients["google"] = oauth.register(
            name="google",
            client_id=google_client_id,
            client_secret=google_client_secret,
            server_metadata_url="https://accounts.google.com/.well-known/openid-configuration",
            client_kwargs={"scope": "openid email profile"},
        )

    if ms_client_id and ms_client_secret:
        oauth_clients["microsoft"] = oauth.register(
            name="microsoft",
            client_id=ms_client_id,
            client_secret=ms_client_secret,
            server_metadata_url="https://login.microsoftonline.com/common/v2.0/.well-known/openid-configuration",
            client_kwargs={"scope": "openid email profile User.Read"},
        )


def get_conn():
    return sqlite3.connect(DB_PATH)


def utc_now_iso() -> str:
    return datetime.now(timezone.utc).replace(microsecond=0).isoformat()

def _sha256_hex(s: str) -> str:
    return hashlib.sha256((s or "").encode("utf-8")).hexdigest()


def _record_auth_attempt(action: str, key: str, ip: str):
    # Best-effort; auth should still work if this logging fails.
    try:
        conn = get_conn()
        cur = conn.cursor()
        cur.execute(
            "INSERT INTO auth_attempts (action, key, ip, created_at_utc) VALUES (?, ?, ?, ?)",
            ((action or "").strip(), (key or "").strip().lower(), (ip or "").strip(), utc_now_iso()),
        )
        conn.commit()
        conn.close()
    except Exception:
        pass


def _rate_limited(action: str, key: str, ip: str, limit: int, window_seconds: int) -> tuple[bool, str]:
    """
    Returns (limited, reset_at_iso). Uses separate buckets for IP and (optional) key.
    """
    limit = max(1, int(limit or 1))
    window_seconds = max(1, int(window_seconds or 1))
    cutoff = (datetime.now(timezone.utc) - timedelta(seconds=window_seconds)).replace(microsecond=0).isoformat()
    key_norm = (key or "").strip().lower()
    try:
        conn = get_conn()
        cur = conn.cursor()
        ip_count = int(
            cur.execute(
                "SELECT COUNT(*) FROM auth_attempts WHERE action = ? AND ip = ? AND created_at_utc >= ?",
                (action, ip, cutoff),
            ).fetchone()[0]
            or 0
        )
        key_count = 0
        if key_norm:
            key_count = int(
                cur.execute(
                    "SELECT COUNT(*) FROM auth_attempts WHERE action = ? AND key = ? AND created_at_utc >= ?",
                    (action, key_norm, cutoff),
                ).fetchone()[0]
                or 0
            )
        conn.close()
        if ip_count >= limit or key_count >= limit:
            reset_at = (datetime.now(timezone.utc) + timedelta(seconds=window_seconds)).replace(microsecond=0).isoformat()
            return True, reset_at
        return False, ""
    except Exception:
        return False, ""


def _password_policy_ok(pw: str) -> tuple[bool, str]:
    pw = pw or ""
    min_len = max(6, _parse_int_env("PASSWORD_MIN_LEN", 8))
    if len(pw) < min_len:
        return False, f"Password must be at least {min_len} characters."
    # Basic complexity without being hostile: require at least 2 of 4 classes.
    classes = 0
    classes += 1 if re.search(r"[a-z]", pw) else 0
    classes += 1 if re.search(r"[A-Z]", pw) else 0
    classes += 1 if re.search(r"[0-9]", pw) else 0
    classes += 1 if re.search(r"[^A-Za-z0-9]", pw) else 0
    if classes < 2:
        return False, "Password must include at least two of: lowercase, uppercase, number, symbol."
    return True, ""


def get_user_brief(user_id: int) -> dict[str, Any] | None:
    conn = get_conn()
    try:
        cur = conn.cursor()
        cur.execute("SELECT id, username, COALESCE(photo_path, '') FROM users WHERE id = ?", (user_id,))
        row = cur.fetchone()
        if not row:
            return None
        return {"id": int(row[0]), "username": row[1], "photo_path": row[2] or ""}
    finally:
        conn.close()


def list_other_users(user_id: int) -> list[dict[str, Any]]:
    conn = get_conn()
    try:
        cur = conn.cursor()
        cur.execute(
            "SELECT id, username, COALESCE(photo_path, '') FROM users WHERE id != ? ORDER BY username COLLATE NOCASE",
            (user_id,),
        )
        return [{"id": int(r[0]), "username": r[1], "photo_path": r[2] or ""} for r in cur.fetchall()]
    finally:
        conn.close()


def _get_user_id_by_username(username: str) -> int | None:
    u = (username or "").strip()
    if not u:
        return None
    conn = get_conn()
    try:
        cur = conn.cursor()
        cur.execute("SELECT id FROM users WHERE username = ? COLLATE NOCASE LIMIT 1", (u,))
        row = cur.fetchone()
        return int(row[0]) if row else None
    finally:
        conn.close()


def follow_counts(user_id: int) -> dict[str, int]:
    conn = get_conn()
    try:
        cur = conn.cursor()
        cur.execute("SELECT COUNT(*) FROM user_follows WHERE followed_id = ?", (user_id,))
        followers = int(cur.fetchone()[0] or 0)
        cur.execute("SELECT COUNT(*) FROM user_follows WHERE follower_id = ?", (user_id,))
        following = int(cur.fetchone()[0] or 0)
        return {"followers": followers, "following": following}
    finally:
        conn.close()


def is_following(viewer_id: int, target_id: int) -> bool:
    if not viewer_id or not target_id or viewer_id == target_id:
        return False
    conn = get_conn()
    try:
        cur = conn.cursor()
        cur.execute(
            "SELECT 1 FROM user_follows WHERE follower_id = ? AND followed_id = ? LIMIT 1",
            (viewer_id, target_id),
        )
        return cur.fetchone() is not None
    finally:
        conn.close()


def set_follow(viewer_id: int, target_id: int, follow: bool) -> None:
    if not viewer_id or not target_id or viewer_id == target_id:
        return
    conn = get_conn()
    try:
        cur = conn.cursor()
        if follow:
            cur.execute(
                "INSERT OR IGNORE INTO user_follows (follower_id, followed_id, created_at_utc) VALUES (?, ?, ?)",
                (viewer_id, target_id, utc_now_iso()),
            )
        else:
            cur.execute(
                "DELETE FROM user_follows WHERE follower_id = ? AND followed_id = ?",
                (viewer_id, target_id),
            )
        conn.commit()
    finally:
        conn.close()


def _escape_like(s: str) -> str:
    """
    Escape user input for SQLite LIKE ... ESCAPE '\\' patterns.
    Treat %, _ and \\ as literals.
    """
    return (s or "").replace("\\", "\\\\").replace("%", "\\%").replace("_", "\\_")


def list_users_for_viewer(viewer_id: int, q: str = "", limit: int = 50) -> list[dict[str, Any]]:
    query = (q or "").strip()
    like = f"%{_escape_like(query)}%"
    conn = get_conn()
    try:
        cur = conn.cursor()
        cur.execute(
            """
            SELECT u.id, u.username, COALESCE(u.photo_path, '') AS photo_path,
                   CASE WHEN f.follower_id IS NULL THEN 0 ELSE 1 END AS is_following
            FROM users u
            LEFT JOIN user_follows f
              ON f.follower_id = ? AND f.followed_id = u.id
            WHERE u.id != ?
              AND (
                ? = ''
                OR (u.username COLLATE NOCASE) LIKE ? ESCAPE '\\'
                OR (COALESCE(u.email, '') COLLATE NOCASE) LIKE ? ESCAPE '\\'
              )
            ORDER BY u.username COLLATE NOCASE
            LIMIT ?
            """,
            (viewer_id, viewer_id, query, like, like, max(1, int(limit))),
        )
        out = []
        for r in cur.fetchall():
            out.append(
                {
                    "id": int(r[0]),
                    "username": r[1],
                    "photo_path": r[2] or "",
                    "is_following": bool(int(r[3] or 0)),
                }
            )
        return out
    finally:
        conn.close()


def ensure_dm_conversation(user_a: int, user_b: int) -> int:
    if user_a == user_b:
        raise ValueError("Cannot create DM with self")
    conn = get_conn()
    try:
        cur = conn.cursor()
        cur.execute(
            """
            SELECT c.id
            FROM chat_conversations c
            JOIN chat_members m ON m.conversation_id = c.id
            WHERE c.is_group = 0 AND m.user_id IN (?, ?)
            GROUP BY c.id
            HAVING COUNT(DISTINCT m.user_id) = 2
               AND (SELECT COUNT(*) FROM chat_members cm WHERE cm.conversation_id = c.id) = 2
            LIMIT 1
            """,
            (user_a, user_b),
        )
        row = cur.fetchone()
        if row:
            return int(row[0])

        created_at = utc_now_iso()
        cur.execute(
            "INSERT INTO chat_conversations (is_group, title, created_at_utc) VALUES (0, NULL, ?)",
            (created_at,),
        )
        conv_id = int(cur.lastrowid)
        joined_at = created_at
        cur.execute(
            "INSERT OR IGNORE INTO chat_members (conversation_id, user_id, role, joined_at_utc) VALUES (?, ?, ?, ?)",
            (conv_id, user_a, "member", joined_at),
        )
        cur.execute(
            "INSERT OR IGNORE INTO chat_members (conversation_id, user_id, role, joined_at_utc) VALUES (?, ?, ?, ?)",
            (conv_id, user_b, "member", joined_at),
        )
        conn.commit()
        return conv_id
    finally:
        conn.close()


def user_is_member(conversation_id: int, user_id: int) -> bool:
    conn = get_conn()
    try:
        cur = conn.cursor()
        cur.execute(
            "SELECT 1 FROM chat_members WHERE conversation_id = ? AND user_id = ? LIMIT 1",
            (conversation_id, user_id),
        )
        return cur.fetchone() is not None
    finally:
        conn.close()


def conversation_member_ids(conversation_id: int) -> list[int]:
    conn = get_conn()
    try:
        cur = conn.cursor()
        cur.execute("SELECT user_id FROM chat_members WHERE conversation_id = ?", (conversation_id,))
        return [int(r[0]) for r in cur.fetchall()]
    finally:
        conn.close()


def conversation_brief_for_user(conversation_id: int, user_id: int) -> dict[str, Any] | None:
    if not user_is_member(conversation_id, user_id):
        return None
    conn = get_conn()
    try:
        cur = conn.cursor()
        cur.execute(
            "SELECT id, is_group, COALESCE(title, ''), created_at_utc FROM chat_conversations WHERE id = ?",
            (conversation_id,),
        )
        c = cur.fetchone()
        if not c:
            return None

        cur.execute(
            """
            SELECT id, sender_id, COALESCE(body, ''), created_at_utc
            FROM chat_messages
            WHERE conversation_id = ?
            ORDER BY id DESC
            LIMIT 1
            """,
            (conversation_id,),
        )
        last = cur.fetchone()

        # DM label: show the other person's username.
        title = c[2] or ""
        is_group = int(c[1]) == 1
        peer = None
        if not is_group:
            cur.execute(
                """
                SELECT u.id, u.username, COALESCE(u.photo_path, '')
                FROM chat_members m
                JOIN users u ON u.id = m.user_id
                WHERE m.conversation_id = ? AND u.id != ?
                LIMIT 1
                """,
                (conversation_id, user_id),
            )
            p = cur.fetchone()
            if p:
                peer = {"id": int(p[0]), "username": p[1], "photo_path": p[2] or ""}
                if not title:
                    title = peer["username"]

        # Unread: count messages not marked as read by this user (excluding own messages).
        cur.execute(
            """
            SELECT COUNT(*)
            FROM chat_messages msg
            WHERE msg.conversation_id = ?
              AND msg.sender_id != ?
              AND msg.id NOT IN (SELECT message_id FROM chat_reads r WHERE r.user_id = ?)
            """,
            (conversation_id, user_id, user_id),
        )
        unread = int(cur.fetchone()[0] or 0)

        return {
            "id": int(c[0]),
            "is_group": bool(is_group),
            "title": title,
            "peer": peer,
            "last": (
                None
                if not last
                else {"id": int(last[0]), "sender_id": int(last[1]), "body": last[2], "created_at_utc": last[3]}
            ),
            "unread": unread,
        }
    finally:
        conn.close()


def list_conversations_for_user(user_id: int) -> list[dict[str, Any]]:
    conn = get_conn()
    try:
        cur = conn.cursor()
        cur.execute(
            """
            SELECT c.id
            FROM chat_conversations c
            JOIN chat_members m ON m.conversation_id = c.id
            WHERE m.user_id = ?
            ORDER BY COALESCE(
                (SELECT created_at_utc FROM chat_messages mm WHERE mm.conversation_id = c.id ORDER BY mm.id DESC LIMIT 1),
                c.created_at_utc
            ) DESC
            """,
            (user_id,),
        )
        ids = [int(r[0]) for r in cur.fetchall()]
    finally:
        conn.close()
    items = []
    for cid in ids:
        brief = conversation_brief_for_user(cid, user_id)
        if brief:
            items.append(brief)
    return items


def list_messages(conversation_id: int, user_id: int, limit: int = 80) -> list[dict[str, Any]]:
    if not user_is_member(conversation_id, user_id):
        return []
    conn = get_conn()
    try:
        cur = conn.cursor()
        cur.execute(
            """
            SELECT id, sender_id, type, COALESCE(body, ''), COALESCE(media_path, ''), created_at_utc
            FROM chat_messages
            WHERE conversation_id = ?
            ORDER BY id DESC
            LIMIT ?
            """,
            (conversation_id, int(max(1, min(limit, 200)))),
        )
        rows = cur.fetchall()
        rows.reverse()

        # For DMs we can show a simple "read by peer" indicator based on presence of chat_reads record.
        member_ids = conversation_member_ids(conversation_id)
        peer_id = next((mid for mid in member_ids if mid != user_id), None)
        read_by_peer_ids: set[int] = set()
        if peer_id is not None:
            cur.execute(
                "SELECT message_id FROM chat_reads WHERE user_id = ?",
                (peer_id,),
            )
            read_by_peer_ids = {int(r[0]) for r in cur.fetchall()}

        return [
            {
                "id": int(r[0]),
                "sender_id": int(r[1]),
                "type": r[2],
                "body": r[3],
                "media_path": r[4],
                "created_at_utc": r[5],
                "read_by_peer": (peer_id is not None and int(r[0]) in read_by_peer_ids),
            }
            for r in rows
        ]
    finally:
        conn.close()


def mark_conversation_read(conversation_id: int, user_id: int) -> None:
    if not user_is_member(conversation_id, user_id):
        return
    conn = get_conn()
    try:
        cur = conn.cursor()
        now = utc_now_iso()
        # Mark all received messages as read.
        cur.execute(
            """
            INSERT OR IGNORE INTO chat_reads (message_id, user_id, read_at_utc)
            SELECT id, ?, ?
            FROM chat_messages
            WHERE conversation_id = ? AND sender_id != ?
            """,
            (user_id, now, conversation_id, user_id),
        )
        conn.commit()
    finally:
        conn.close()


DEFAULT_AI_MODEL = (os.getenv("OPENAI_MODEL") or "gpt-4.1-mini").strip() or "gpt-4.1-mini"
DEFAULT_AI_VISION_MODEL = (os.getenv("OPENAI_VISION_MODEL") or "").strip()
AI_HISTORY_LIMIT = max(10, _parse_int_env("AI_HISTORY_LIMIT", 40))
AI_MEMORY_LIMIT = max(0, _parse_int_env("AI_MEMORY_LIMIT", 12))
AI_FILE_MAX_BYTES = max(256 * 1024, _parse_int_env("AI_FILE_MAX_BYTES", 8 * 1024 * 1024))


def _safe_now_iso() -> str:
    return utc_now_iso()


def _public_base_url() -> str:
    # Best-effort; use PUBLIC_BASE_URL in prod to create correct share links behind proxies.
    return (os.getenv("PUBLIC_BASE_URL") or "").strip()


def _guess_lan_ip() -> str | None:
    """
    Best-effort local/LAN IP detection for development share links.
    Returns None if we can't find a non-loopback IPv4.
    """
    # 1) UDP connect trick (does not need to send traffic successfully).
    try:
        s = socket.socket(socket.AF_INET, socket.SOCK_DGRAM)
        try:
            s.connect(("8.8.8.8", 80))
            ip = str(s.getsockname()[0] or "").strip()
        finally:
            try:
                s.close()
            except Exception:
                pass
        if ip and not ip.startswith("127."):
            return ip
    except Exception:
        pass

    # 2) Hostname resolution fallback (often returns 127.0.0.1; ignore that).
    try:
        ip = str(socket.gethostbyname(socket.gethostname()) or "").strip()
        if ip and not ip.startswith("127."):
            return ip
    except Exception:
        pass

    return None


def _best_public_base_url() -> str:
    """
    Returns a base URL suitable for share/billing redirects.

    Priority:
    - PUBLIC_BASE_URL env var (recommended for prod / ngrok)
    - request.url_root (current request host)
    - if current host is localhost/127.0.0.1: try to replace it with LAN IP
    """
    base = _public_base_url().rstrip("/")
    if base:
        return base

    try:
        base = (request.url_root or "").rstrip("/")
    except Exception:
        base = ""

    if not base:
        return ""

    try:
        u = urlsplit(base)
        host = (u.hostname or "").strip().lower()
        if host in ("127.0.0.1", "localhost"):
            ip = _guess_lan_ip()
            if ip:
                netloc = ip
                if u.port:
                    netloc = f"{ip}:{u.port}"
                return urlunsplit((u.scheme, netloc, "", "", ""))
    except Exception:
        pass

    return base


def _stripe_enabled() -> bool:
    return stripe is not None and bool((os.getenv("STRIPE_SECRET_KEY") or "").strip())


def _stripe_price_ids() -> dict[str, dict[str, str]]:
    """
    Returns mapping: plan -> {month|year: price_id}.
    Configure via env:
      STRIPE_PRICE_PRO_MONTHLY, STRIPE_PRICE_PRO_YEARLY,
      STRIPE_PRICE_ULTIMATE_MONTHLY, STRIPE_PRICE_ULTIMATE_YEARLY
    """
    return {
        "pro": {
            "month": (os.getenv("STRIPE_PRICE_PRO_MONTHLY") or "").strip(),
            "year": (os.getenv("STRIPE_PRICE_PRO_YEARLY") or "").strip(),
        },
        "ultimate": {
            "month": (os.getenv("STRIPE_PRICE_ULTIMATE_MONTHLY") or "").strip(),
            "year": (os.getenv("STRIPE_PRICE_ULTIMATE_YEARLY") or "").strip(),
        },
    }


def _stripe_set_api_key() -> None:
    if stripe is None:
        return
    stripe.api_key = (os.getenv("STRIPE_SECRET_KEY") or "").strip()


def _stripe_webhook_secret() -> str:
    return (os.getenv("STRIPE_WEBHOOK_SECRET") or "").strip()


def _stripe_price_id_to_plan(price_id: str) -> str:
    pid = (price_id or "").strip()
    if not pid:
        return ""
    prices = _stripe_price_ids()
    for plan, by_int in (prices or {}).items():
        for _intv, v in (by_int or {}).items():
            if v and str(v).strip() == pid:
                return str(plan).strip().lower()
    return ""


def _stripe_ts_to_iso(ts: int | float | None) -> str:
    try:
        t = int(ts or 0)
    except Exception:
        t = 0
    if t <= 0:
        return ""
    return datetime.fromtimestamp(t, tz=timezone.utc).isoformat()


def _stripe_get_customer_metadata(customer_id: str) -> dict[str, str]:
    if not customer_id:
        return {}
    try:
        c = stripe.Customer.retrieve(customer_id)  # type: ignore[union-attr]
        md = c.get("metadata") if isinstance(c, dict) else {}
        return {str(k): str(v) for k, v in (md or {}).items() if k and v}
    except Exception:
        return {}


def _stripe_resolve_user(firebase_uid: str, local_user_id: str | None) -> tuple[str, int | None]:
    uid = (firebase_uid or "").strip()
    uid = uid.strip()
    user_id: int | None = None
    if local_user_id and str(local_user_id).strip():
        try:
            user_id = int(str(local_user_id).strip())
        except ValueError:
            user_id = None
    if user_id is None and uid:
        user_id = _user_id_for_firebase_uid(uid)
    return uid, user_id


def _stripe_apply_subscription_state(
    *,
    firebase_uid: str,
    user_id: int | None,
    customer_id: str,
    subscription_id: str,
    status: str,
    current_period_end_utc: str,
    cancel_at_period_end: bool,
    price_id: str,
    interval: str,
    plan: str,
) -> None:
    # Decide entitlements: paid only when Stripe says active/trialing.
    st = (status or "").strip().lower()
    paid = st in {"active", "trialing"}
    desired_plan = (plan or "").strip().lower() if paid else FREE_PLAN
    if desired_plan not in {FREE_PLAN, *PRO_PLANS}:
        desired_plan = FREE_PLAN

    if user_id is not None:
        try:
            _set_user_plan(int(user_id), desired_plan)
        except Exception:
            pass
        try:
            if customer_id:
                _user_set_stripe_customer_id(int(user_id), customer_id)
        except Exception:
            pass
        try:
            _billing_upsert_subscription(
                int(user_id),
                customer_id=customer_id,
                subscription_id=subscription_id,
                price_id=price_id,
                status=st,
                current_period_end_utc=current_period_end_utc,
                cancel_at_period_end=bool(cancel_at_period_end),
                plan=(plan or "").strip().lower(),
                interval=(interval or "").strip().lower(),
            )
        except Exception:
            pass

    # Firestore is authoritative for plan/usage in prod. Best-effort update.
    uid = (firebase_uid or "").strip()
    if uid:
        try:
            doc_ref = _billing_firestore_doc(uid)
            if doc_ref is not None:
                doc_ref.set(
                    {
                        "uid": uid,
                        "plan_type": desired_plan,
                        "subscription_status": st,
                        "current_period_end_utc": current_period_end_utc,
                        "cancel_at_period_end": bool(cancel_at_period_end),
                        "stripe_customer_id": customer_id or "",
                        "stripe_subscription_id": subscription_id or "",
                        "stripe_price_id": price_id or "",
                        "stripe_interval": (interval or "").strip().lower(),
                        "updated_at_utc": _billing_firestore_now_iso(),
                    },
                    merge=True,
                )
        except Exception:
            pass


def _stripe_sync_subscription(subscription_id: str, customer_id: str, *, firebase_uid: str = "", local_user_id: str | None = None) -> None:
    """
    Retrieve a subscription from Stripe and update entitlements in SQLite/Firestore.
    """
    if not subscription_id:
        return
    try:
        sub = stripe.Subscription.retrieve(subscription_id, expand=["items.data.price"])  # type: ignore[union-attr]
    except Exception:
        return
    if not isinstance(sub, dict):
        return
    md = sub.get("metadata") if isinstance(sub.get("metadata"), dict) else {}
    uid = str(md.get("firebase_uid") or firebase_uid or "").strip()
    luid = str(md.get("local_user_id") or local_user_id or "").strip() or None
    if not uid and customer_id:
        uid = str(_stripe_get_customer_metadata(customer_id).get("firebase_uid") or "").strip()
    uid, user_id = _stripe_resolve_user(uid, luid)

    st = str(sub.get("status") or "").strip().lower()
    cancel_at_period_end = bool(sub.get("cancel_at_period_end") or False)
    current_period_end_utc = _stripe_ts_to_iso(sub.get("current_period_end"))

    items = ((sub.get("items") or {}).get("data") or []) if isinstance(sub.get("items"), dict) else []
    price_obj = (items[0].get("price") if items and isinstance(items[0], dict) else {}) if items else {}
    price_id = str((price_obj or {}).get("id") or "").strip()
    recurring = (price_obj or {}).get("recurring") if isinstance((price_obj or {}).get("recurring"), dict) else {}
    interval = str((recurring or {}).get("interval") or "").strip().lower()
    plan = _stripe_price_id_to_plan(price_id) or str((md or {}).get("plan") or "pro").strip().lower()
    if plan not in {FREE_PLAN, *PRO_PLANS}:
        plan = "pro"

    _stripe_apply_subscription_state(
        firebase_uid=uid,
        user_id=user_id,
        customer_id=str(customer_id or sub.get("customer") or "").strip(),
        subscription_id=str(subscription_id).strip(),
        status=st,
        current_period_end_utc=current_period_end_utc,
        cancel_at_period_end=cancel_at_period_end,
        price_id=price_id,
        interval=interval,
        plan=plan,
    )


def _stripe_sync_subscription_object(sub: dict[str, Any], customer_id: str, *, firebase_uid: str = "", local_user_id: str | None = None) -> None:
    """
    Sync using the subscription object from a webhook event (avoids an extra API call and
    still works for deleted subscriptions).
    """
    if not isinstance(sub, dict):
        return
    md = sub.get("metadata") if isinstance(sub.get("metadata"), dict) else {}
    uid = str((md or {}).get("firebase_uid") or firebase_uid or "").strip()
    luid = str((md or {}).get("local_user_id") or local_user_id or "").strip() or None
    if not uid and customer_id:
        uid = str(_stripe_get_customer_metadata(customer_id).get("firebase_uid") or "").strip()
    uid, user_id = _stripe_resolve_user(uid, luid)

    st = str(sub.get("status") or "").strip().lower()
    cancel_at_period_end = bool(sub.get("cancel_at_period_end") or False)
    current_period_end_utc = _stripe_ts_to_iso(sub.get("current_period_end"))

    items = ((sub.get("items") or {}).get("data") or []) if isinstance(sub.get("items"), dict) else []
    price_obj = (items[0].get("price") if items and isinstance(items[0], dict) else {}) if items else {}
    price_id = str((price_obj or {}).get("id") or "").strip()
    recurring = (price_obj or {}).get("recurring") if isinstance((price_obj or {}).get("recurring"), dict) else {}
    interval = str((recurring or {}).get("interval") or "").strip().lower()
    plan = _stripe_price_id_to_plan(price_id) or str((md or {}).get("plan") or "pro").strip().lower()
    if plan not in {FREE_PLAN, *PRO_PLANS}:
        plan = "pro"

    _stripe_apply_subscription_state(
        firebase_uid=uid,
        user_id=user_id,
        customer_id=str(customer_id or sub.get("customer") or "").strip(),
        subscription_id=str(sub.get("id") or "").strip(),
        status=st,
        current_period_end_utc=current_period_end_utc,
        cancel_at_period_end=cancel_at_period_end,
        price_id=price_id,
        interval=interval,
        plan=plan,
    )


def _billing_upsert_subscription(
    user_id: int,
    *,
    customer_id: str = "",
    subscription_id: str = "",
    price_id: str = "",
    status: str = "",
    current_period_end_utc: str = "",
    cancel_at_period_end: bool = False,
    plan: str = "",
    interval: str = "",
) -> None:
    now = utc_now_iso()
    conn = get_conn()
    try:
        cur = conn.cursor()
        cur.execute(
            """
            INSERT INTO billing_subscriptions (
              user_id, provider, customer_id, subscription_id, price_id, status, current_period_end_utc,
              cancel_at_period_end, plan, interval, updated_at_utc
            )
            VALUES (?, 'stripe', ?, ?, ?, ?, ?, ?, ?, ?, ?)
            ON CONFLICT(user_id) DO UPDATE SET
              customer_id=excluded.customer_id,
              subscription_id=excluded.subscription_id,
              price_id=excluded.price_id,
              status=excluded.status,
              current_period_end_utc=excluded.current_period_end_utc,
              cancel_at_period_end=excluded.cancel_at_period_end,
              plan=excluded.plan,
              interval=excluded.interval,
              updated_at_utc=excluded.updated_at_utc
            """,
            (
                user_id,
                customer_id or "",
                subscription_id or "",
                price_id or "",
                status or "",
                current_period_end_utc or "",
                1 if cancel_at_period_end else 0,
                plan or "",
                interval or "",
                now,
            ),
        )
        conn.commit()
    finally:
        conn.close()


def _billing_get_subscription(user_id: int) -> dict[str, Any] | None:
    conn = get_conn()
    try:
        cur = conn.cursor()
        cur.execute(
            """
            SELECT customer_id, subscription_id, price_id, status, current_period_end_utc, cancel_at_period_end, plan, interval, updated_at_utc
            FROM billing_subscriptions
            WHERE user_id = ?
            """,
            (user_id,),
        )
        r = cur.fetchone()
        if not r:
            return None
        return {
            "customer_id": r[0] or "",
            "subscription_id": r[1] or "",
            "price_id": r[2] or "",
            "status": r[3] or "",
            "current_period_end_utc": r[4] or "",
            "cancel_at_period_end": bool(int(r[5] or 0)),
            "plan": r[6] or "",
            "interval": r[7] or "",
            "updated_at_utc": r[8] or "",
        }
    finally:
        conn.close()


def _user_set_stripe_customer_id(user_id: int, customer_id: str) -> None:
    conn = get_conn()
    try:
        cur = conn.cursor()
        cur.execute("UPDATE users SET stripe_customer_id = ? WHERE id = ?", (customer_id or "", user_id))
        conn.commit()
    finally:
        conn.close()


def _user_get_stripe_customer_id(user_id: int) -> str:
    conn = get_conn()
    try:
        cur = conn.cursor()
        cur.execute("SELECT COALESCE(stripe_customer_id, '') FROM users WHERE id = ?", (user_id,))
        row = cur.fetchone()
        return (row[0] or "") if row else ""
    finally:
        conn.close()


def _set_user_plan(user_id: int, plan: str) -> None:
    p = (plan or "").strip().lower()
    if p not in {FREE_PLAN, *PRO_PLANS}:
        p = FREE_PLAN
    conn = get_conn()
    try:
        cur = conn.cursor()
        cur.execute("UPDATE users SET plan = ? WHERE id = ?", (p, user_id))
        conn.commit()
    finally:
        conn.close()


# -----------------------------
# Billing/Usage Store (Firestore preferred; SQLite fallback)
# -----------------------------

FIRESTORE_BILLING_COLLECTION = (os.getenv("FIRESTORE_BILLING_COLLECTION") or "billing_users").strip() or "billing_users"


def _billing_firestore_client():
    """
    Returns a Firestore client or None if Firestore is not available/configured.
    We intentionally treat Firestore as optional so local dev keeps working.
    """
    if fb_firestore is None or gc_firestore is None:
        return None
    fb_err = _firebase_init_error()
    if fb_err:
        return None
    try:
        return fb_firestore.client()  # type: ignore[union-attr]
    except Exception:
        return None


def _billing_firestore_doc(uid: str):
    client = _billing_firestore_client()
    if client is None:
        return None
    return client.collection(FIRESTORE_BILLING_COLLECTION).document(uid)


def _billing_firestore_now_iso() -> str:
    return datetime.now(timezone.utc).isoformat()


def _billing_firestore_ensure_user(uid: str, email: str, *, user_id: int | None = None, plan_hint: str = FREE_PLAN) -> dict[str, Any] | None:
    """
    Ensure a billing doc exists. Returns the doc dict (best-effort).
    """
    uid = (uid or "").strip()
    if not uid:
        return None
    doc = _billing_firestore_doc(uid)
    if doc is None:
        return None
    now = _billing_firestore_now_iso()
    try:
        snap = doc.get()
        base = {
            "uid": uid,
            "email": (email or "").strip().lower(),
            "plan_type": (plan_hint or FREE_PLAN).strip().lower() or FREE_PLAN,
            "messages_used_today": 0,
            "last_reset_date": "",
            "subscription_status": "",
            "stripe_customer_id": "",
            "stripe_subscription_id": "",
            "current_period_end_utc": "",
            "cancel_at_period_end": False,
            "created_at_utc": now,
            "updated_at_utc": now,
        }
        if user_id is not None:
            base["local_user_id"] = int(user_id)
        if not snap.exists:
            doc.set(base, merge=True)
            return base
        existing = snap.to_dict() or {}
        # Keep identity fresh.
        updates: dict[str, Any] = {"updated_at_utc": now}
        if email and (str(existing.get("email") or "").strip().lower() != email.strip().lower()):
            updates["email"] = email.strip().lower()
        if user_id is not None and (existing.get("local_user_id") is None):
            updates["local_user_id"] = int(user_id)
        # First write: if plan_type missing, seed with hint.
        if not str(existing.get("plan_type") or "").strip():
            updates["plan_type"] = (plan_hint or FREE_PLAN).strip().lower() or FREE_PLAN
        if len(updates) > 1:
            doc.set(updates, merge=True)
        return {**base, **existing, **updates}
    except Exception:
        return None


def _billing_firestore_effective_plan(doc_dict: dict[str, Any] | None, plan_hint: str) -> str:
    """
    Compute effective plan from the Firestore doc.
    - If subscription_status is active/trialing -> keep paid plan_type (usually pro/ultimate).
    - Otherwise, only treat plan_type as paid if it was set manually (no Stripe ids).
    """
    if not doc_dict:
        p = (plan_hint or FREE_PLAN).strip().lower()
        return p if p in {FREE_PLAN, *PRO_PLANS} else FREE_PLAN

    p = str(doc_dict.get("plan_type") or plan_hint or FREE_PLAN).strip().lower() or FREE_PLAN
    if p not in {FREE_PLAN, *PRO_PLANS}:
        p = FREE_PLAN

    status = str(doc_dict.get("subscription_status") or "").strip().lower()
    sub_id = str(doc_dict.get("stripe_subscription_id") or "").strip()
    cust_id = str(doc_dict.get("stripe_customer_id") or "").strip()
    has_stripe = bool(sub_id or cust_id)

    if p in PRO_PLANS and has_stripe:
        if status in {"active", "trialing"}:
            return p
        return FREE_PLAN

    return p


def _billing_firestore_reserve_daily_message(
    uid: str,
    email: str,
    *,
    plan_hint: str,
    requested_tz: str | None,
    free_daily_limit: int,
    user_id: int | None = None,
) -> tuple[bool, dict[str, Any], bool]:
    """
    Firestore-backed reserve.
    Returns (ok, quota_payload, reserved).
    """
    client = _billing_firestore_client()
    if client is None:
        return True, {"plan": plan_hint}, False
    uid = (uid or "").strip()
    if not uid:
        return True, {"plan": plan_hint}, False

    # Keep same TZ rules as SQLite store.
    user_tz = normalize_tz_name(str(requested_tz)) or normalize_tz_name((os.getenv("USAGE_TZ") or "").strip())
    day = usage_day_iso(user_tz)
    limit = int(max(0, free_daily_limit))

    doc_ref = client.collection(FIRESTORE_BILLING_COLLECTION).document(uid)
    now = _billing_firestore_now_iso()

    # Quota disabled: still ensure identity exists, but don't reserve.
    if limit <= 0:
        _billing_firestore_ensure_user(uid, email, user_id=user_id, plan_hint=plan_hint)
        return True, {"plan": FREE_PLAN, "store": "firestore", "day": day, "tz": user_tz or ""}, False

    tx = client.transaction()

    transactional = getattr(gc_firestore, "transactional", None) if gc_firestore is not None else None
    if transactional is None:
        return True, {"plan": plan_hint}, False

    @transactional  # type: ignore[misc]
    def _run(transaction):
        snap = doc_ref.get(transaction=transaction)
        d = snap.to_dict() if snap.exists else {}

        # Seed doc if missing.
        if not snap.exists:
            seed = {
                "uid": uid,
                "email": (email or "").strip().lower(),
                "plan_type": (plan_hint or FREE_PLAN).strip().lower() or FREE_PLAN,
                "messages_used_today": 0,
                "last_reset_date": "",
                "subscription_status": "",
                "stripe_customer_id": "",
                "stripe_subscription_id": "",
                "current_period_end_utc": "",
                "cancel_at_period_end": False,
                "created_at_utc": now,
                "updated_at_utc": now,
            }
            if user_id is not None:
                seed["local_user_id"] = int(user_id)
            transaction.set(doc_ref, seed, merge=True)
            d = seed

        effective_plan = _billing_firestore_effective_plan(d, plan_hint)
        if effective_plan in PRO_PLANS:
            transaction.set(doc_ref, {"updated_at_utc": now}, merge=True)
            return True, {"plan": effective_plan, "store": "firestore"}, False

        last_reset = str(d.get("last_reset_date") or "").strip()
        used = int(d.get("messages_used_today") or 0)
        if last_reset != day:
            used = 0

        if used >= limit:
            transaction.set(
                doc_ref,
                {
                    "plan_type": FREE_PLAN,
                    "last_reset_date": day,
                    "messages_used_today": used,
                    "updated_at_utc": now,
                },
                merge=True,
            )
            return (
                False,
                {
                    "plan": FREE_PLAN,
                    "store": "firestore",
                    "limit": limit,
                    "used": used,
                    "day": day,
                    "tz": user_tz or "",
                    "reset_at": next_reset_iso(user_tz),
                },
                False,
            )

        used_after = used + 1
        transaction.set(
            doc_ref,
            {
                "plan_type": FREE_PLAN,
                "email": (email or "").strip().lower(),
                "last_reset_date": day,
                "messages_used_today": used_after,
                "updated_at_utc": now,
            },
            merge=True,
        )
        return (
            True,
            {
                "plan": FREE_PLAN,
                "store": "firestore",
                "limit": limit,
                "used": used_after,
                "day": day,
                "tz": user_tz or "",
                "reset_at": next_reset_iso(user_tz),
            },
            True,
        )

    try:
        return _run(tx)
    except Exception:
        # Fail open to avoid breaking chat if Firestore has transient issues.
        return True, {"plan": plan_hint}, False


def _billing_firestore_release_daily_message(uid: str, day: str) -> None:
    client = _billing_firestore_client()
    transactional = getattr(gc_firestore, "transactional", None) if gc_firestore is not None else None
    if client is None or transactional is None:
        return
    uid = (uid or "").strip()
    if not uid or not (day or "").strip():
        return
    doc_ref = client.collection(FIRESTORE_BILLING_COLLECTION).document(uid)
    now = _billing_firestore_now_iso()
    tx = client.transaction()

    @transactional  # type: ignore[misc]
    def _run(transaction):
        snap = doc_ref.get(transaction=transaction)
        if not snap.exists:
            return
        d = snap.to_dict() or {}
        if str(d.get("last_reset_date") or "").strip() != str(day).strip():
            return
        used = int(d.get("messages_used_today") or 0)
        if used <= 0:
            return
        transaction.set(doc_ref, {"messages_used_today": used - 1, "updated_at_utc": now}, merge=True)

    try:
        _run(tx)
    except Exception:
        return


def _user_id_for_firebase_uid(uid: str) -> int | None:
    uid = (uid or "").strip()
    if not uid:
        return None
    conn = get_conn()
    try:
        cur = conn.cursor()
        cur.execute(
            "SELECT user_id FROM identities WHERE provider = ? AND provider_user_id = ? LIMIT 1",
            ("firebase", uid),
        )
        row = cur.fetchone()
        return int(row[0]) if row else None
    finally:
        conn.close()


def _ensure_user_dir(*parts: str) -> str:
    path = os.path.join(UPLOAD_FOLDER, *parts)
    os.makedirs(path, exist_ok=True)
    return path


def list_user_memories(user_id: int, limit: int = 20) -> list[dict[str, Any]]:
    if limit <= 0:
        return []
    conn = get_conn()
    try:
        cur = conn.cursor()
        cur.execute(
            """
            SELECT id, text, created_at_utc, COALESCE(source_conversation_id, 0)
            FROM user_memories
            WHERE user_id = ?
            ORDER BY id DESC
            LIMIT ?
            """,
            (user_id, int(max(1, min(limit, 200)))),
        )
        return [
            {"id": int(r[0]), "text": r[1], "created_at_utc": r[2], "source_conversation_id": int(r[3] or 0)}
            for r in cur.fetchall()
        ]
    finally:
        conn.close()


def add_user_memory(user_id: int, text: str, source_conversation_id: int | None = None) -> dict[str, Any] | None:
    t = (text or "").strip()
    if not t:
        return None
    now = _safe_now_iso()
    conn = get_conn()
    try:
        cur = conn.cursor()
        cur.execute(
            "INSERT INTO user_memories (user_id, text, created_at_utc, source_conversation_id) VALUES (?, ?, ?, ?)",
            (user_id, t, now, (int(source_conversation_id) if source_conversation_id else None)),
        )
        mid = int(cur.lastrowid)
        conn.commit()
        return {"id": mid, "text": t, "created_at_utc": now, "source_conversation_id": int(source_conversation_id or 0)}
    finally:
        conn.close()


def delete_user_memory(user_id: int, memory_id: int) -> bool:
    conn = get_conn()
    try:
        cur = conn.cursor()
        cur.execute("DELETE FROM user_memories WHERE id = ? AND user_id = ?", (int(memory_id), user_id))
        conn.commit()
        return cur.rowcount > 0
    finally:
        conn.close()


def get_ai_notes(conversation_id: int, user_id: int) -> dict[str, Any] | None:
    if not _ai_conv_belongs_to_user(conversation_id, user_id):
        return None
    conn = get_conn()
    try:
        cur = conn.cursor()
        cur.execute(
            "SELECT COALESCE(body, ''), updated_at_utc FROM ai_notes WHERE conversation_id = ? AND user_id = ?",
            (conversation_id, user_id),
        )
        row = cur.fetchone()
        if not row:
            return {"body": "", "updated_at_utc": ""}
        return {"body": row[0] or "", "updated_at_utc": row[1] or ""}
    finally:
        conn.close()


def upsert_ai_notes(conversation_id: int, user_id: int, body: str) -> dict[str, Any] | None:
    if not _ai_conv_belongs_to_user(conversation_id, user_id):
        return None
    now = _safe_now_iso()
    conn = get_conn()
    try:
        cur = conn.cursor()
        cur.execute(
            """
            INSERT INTO ai_notes (conversation_id, user_id, body, updated_at_utc)
            VALUES (?, ?, ?, ?)
            ON CONFLICT(conversation_id, user_id) DO UPDATE SET body=excluded.body, updated_at_utc=excluded.updated_at_utc
            """,
            (conversation_id, user_id, (body or ""), now),
        )
        conn.commit()
    finally:
        conn.close()
    return {"body": (body or ""), "updated_at_utc": now}


def list_ai_prompts(user_id: int, limit: int = 60) -> list[dict[str, Any]]:
    conn = get_conn()
    try:
        cur = conn.cursor()
        cur.execute(
            """
            SELECT id, title, body, created_at_utc, updated_at_utc
            FROM ai_prompts
            WHERE user_id = ?
            ORDER BY updated_at_utc DESC
            LIMIT ?
            """,
            (user_id, int(max(1, min(limit, 200)))),
        )
        return [
            {"id": int(r[0]), "title": r[1], "body": r[2], "created_at_utc": r[3], "updated_at_utc": r[4]}
            for r in cur.fetchall()
        ]
    finally:
        conn.close()


def upsert_ai_prompt(user_id: int, title: str, body: str, prompt_id: int | None = None) -> dict[str, Any] | None:
    t = (title or "").strip()
    b = (body or "").strip()
    if not t or not b:
        return None
    now = _safe_now_iso()
    conn = get_conn()
    try:
        cur = conn.cursor()
        if prompt_id:
            cur.execute(
                "UPDATE ai_prompts SET title = ?, body = ?, updated_at_utc = ? WHERE id = ? AND user_id = ?",
                (t, b, now, int(prompt_id), user_id),
            )
            if cur.rowcount <= 0:
                return None
            pid = int(prompt_id)
        else:
            cur.execute(
                "INSERT INTO ai_prompts (user_id, title, body, created_at_utc, updated_at_utc) VALUES (?, ?, ?, ?, ?)",
                (user_id, t, b, now, now),
            )
            pid = int(cur.lastrowid)
        conn.commit()
        return {"id": pid, "title": t, "body": b, "created_at_utc": now, "updated_at_utc": now}
    finally:
        conn.close()


def delete_ai_prompt(user_id: int, prompt_id: int) -> bool:
    conn = get_conn()
    try:
        cur = conn.cursor()
        cur.execute("DELETE FROM ai_prompts WHERE id = ? AND user_id = ?", (int(prompt_id), user_id))
        conn.commit()
        return cur.rowcount > 0
    finally:
        conn.close()


def ai_get_share(conversation_id: int, user_id: int) -> dict[str, Any] | None:
    if not _ai_conv_belongs_to_user(conversation_id, user_id):
        return None
    conn = get_conn()
    try:
        cur = conn.cursor()
        cur.execute(
            """
            SELECT token, created_at_utc, revoked_at_utc
            FROM ai_shares
            WHERE conversation_id = ? AND user_id = ?
            ORDER BY id DESC
            LIMIT 1
            """,
            (conversation_id, user_id),
        )
        row = cur.fetchone()
        if not row:
            return {"token": "", "created_at_utc": "", "revoked_at_utc": ""}
        return {"token": row[0] or "", "created_at_utc": row[1] or "", "revoked_at_utc": row[2] or ""}
    finally:
        conn.close()


def ai_create_or_rotate_share(conversation_id: int, user_id: int) -> dict[str, Any] | None:
    if not _ai_conv_belongs_to_user(conversation_id, user_id):
        return None
    token = uuid4().hex
    now = _safe_now_iso()
    conn = get_conn()
    try:
        cur = conn.cursor()
        # Revoke existing
        cur.execute(
            "UPDATE ai_shares SET revoked_at_utc = ? WHERE conversation_id = ? AND user_id = ? AND (revoked_at_utc IS NULL OR revoked_at_utc = '')",
            (now, conversation_id, user_id),
        )
        cur.execute(
            "INSERT INTO ai_shares (conversation_id, user_id, token, created_at_utc, revoked_at_utc) VALUES (?, ?, ?, ?, '')",
            (conversation_id, user_id, token, now),
        )
        conn.commit()
    finally:
        conn.close()
    return {"token": token, "created_at_utc": now, "revoked_at_utc": ""}


def ai_revoke_share(conversation_id: int, user_id: int) -> bool:
    if not _ai_conv_belongs_to_user(conversation_id, user_id):
        return False
    now = _safe_now_iso()
    conn = get_conn()
    try:
        cur = conn.cursor()
        cur.execute(
            "UPDATE ai_shares SET revoked_at_utc = ? WHERE conversation_id = ? AND user_id = ? AND (revoked_at_utc IS NULL OR revoked_at_utc = '')",
            (now, conversation_id, user_id),
        )
        conn.commit()
        return cur.rowcount > 0
    finally:
        conn.close()


def _ai_conv_belongs_to_user(conversation_id: int, user_id: int) -> bool:
    conn = get_conn()
    try:
        cur = conn.cursor()
        cur.execute("SELECT 1 FROM ai_conversations WHERE id = ? AND user_id = ? LIMIT 1", (conversation_id, user_id))
        return cur.fetchone() is not None
    finally:
        conn.close()


def ai_create_conversation(user_id: int, title: str | None = None, mode: str = "normal", system_prompt: str = "") -> int:
    title = (title or "").strip() or "New chat"
    mode = (mode or "").strip().lower()
    if mode not in CHAT_MODES:
        mode = "normal"
    created = utc_now_iso()
    conn = get_conn()
    try:
        cur = conn.cursor()
        cur.execute(
            """
            INSERT INTO ai_conversations (user_id, title, mode, system_prompt, is_pinned, is_bookmarked, created_at_utc, updated_at_utc)
            VALUES (?, ?, ?, ?, 0, 0, ?, ?)
            """,
            (user_id, title, mode, (system_prompt or "").strip(), created, created),
        )
        conv_id = int(cur.lastrowid)
        conn.commit()
        return conv_id
    finally:
        conn.close()


def ai_conversation_brief(conversation_id: int, user_id: int) -> dict[str, Any] | None:
    conn = get_conn()
    try:
        cur = conn.cursor()
        cur.execute(
            """
            SELECT id, title, mode, system_prompt, is_pinned, is_bookmarked, COALESCE(summary, ''), created_at_utc, updated_at_utc
            FROM ai_conversations
            WHERE id = ? AND user_id = ?
            """,
            (conversation_id, user_id),
        )
        row = cur.fetchone()
        if not row:
            return None
        cur.execute(
            """
            SELECT role, substr(content, 1, 240), created_at_utc
            FROM ai_messages
            WHERE conversation_id = ?
            ORDER BY id DESC
            LIMIT 1
            """,
            (conversation_id,),
        )
        last = cur.fetchone()
        return {
            "id": int(row[0]),
            "title": row[1],
            "mode": row[2],
            "system_prompt": row[3] or "",
            "is_pinned": bool(int(row[4] or 0)),
            "is_bookmarked": bool(int(row[5] or 0)),
            "summary": row[6] or "",
            "created_at_utc": row[7],
            "updated_at_utc": row[8],
            "last": None if not last else {"role": last[0], "preview": (last[1] or ""), "created_at_utc": last[2]},
        }
    finally:
        conn.close()


def ai_list_conversations(user_id: int, q: str = "", limit: int = 50) -> list[dict[str, Any]]:
    q = (q or "").strip()
    limit = int(max(1, min(limit, 200)))
    conn = get_conn()
    try:
        cur = conn.cursor()
        if q:
            like = f"%{q}%"
            cur.execute(
                """
                SELECT DISTINCT c.id
                FROM ai_conversations c
                LEFT JOIN ai_messages m ON m.conversation_id = c.id
                WHERE c.user_id = ?
                  AND (c.title LIKE ? OR m.content LIKE ?)
                ORDER BY c.is_pinned DESC, c.updated_at_utc DESC
                LIMIT ?
                """,
                (user_id, like, like, limit),
            )
        else:
            cur.execute(
                """
                SELECT id
                FROM ai_conversations
                WHERE user_id = ?
                ORDER BY is_pinned DESC, updated_at_utc DESC
                LIMIT ?
                """,
                (user_id, limit),
            )
        ids = [int(r[0]) for r in cur.fetchall()]
    finally:
        conn.close()
    items: list[dict[str, Any]] = []
    for cid in ids:
        b = ai_conversation_brief(cid, user_id)
        if b:
            items.append(b)
    return items


def ai_list_messages(conversation_id: int, user_id: int, limit: int = 200) -> list[dict[str, Any]]:
    if not _ai_conv_belongs_to_user(conversation_id, user_id):
        return []
    limit = int(max(1, min(limit, 500)))
    conn = get_conn()
    try:
        cur = conn.cursor()
        cur.execute(
            """
            SELECT id, role, content, created_at_utc, COALESCE(edited_at_utc, '')
            FROM ai_messages
            WHERE conversation_id = ?
            ORDER BY id ASC
            LIMIT ?
            """,
            (conversation_id, limit),
        )
        return [
            {
                "id": int(r[0]),
                "conversation_id": conversation_id,
                "role": r[1],
                "content": r[2],
                "created_at_utc": r[3],
                "edited_at_utc": (r[4] or ""),
            }
            for r in cur.fetchall()
        ]
    finally:
        conn.close()


def ai_append_message(conversation_id: int, user_id: int, role: str, content: str) -> int:
    if not _ai_conv_belongs_to_user(conversation_id, user_id):
        raise ValueError("Forbidden")
    role = (role or "").strip().lower()
    if role not in {"user", "assistant", "system"}:
        raise ValueError("Invalid role")
    now = utc_now_iso()
    conn = get_conn()
    try:
        cur = conn.cursor()
        cur.execute(
            """
            INSERT INTO ai_messages (conversation_id, role, content, created_at_utc, edited_at_utc)
            VALUES (?, ?, ?, ?, '')
            """,
            (conversation_id, role, content, now),
        )
        mid = int(cur.lastrowid)
        cur.execute("UPDATE ai_conversations SET updated_at_utc = ? WHERE id = ? AND user_id = ?", (now, conversation_id, user_id))
        conn.commit()
        return mid
    finally:
        conn.close()


def ai_delete_message_and_after(conversation_id: int, user_id: int, from_message_id: int) -> None:
    if not _ai_conv_belongs_to_user(conversation_id, user_id):
        return
    conn = get_conn()
    try:
        cur = conn.cursor()
        cur.execute("DELETE FROM ai_messages WHERE conversation_id = ? AND id >= ?", (conversation_id, int(from_message_id)))
        cur.execute(
            "UPDATE ai_conversations SET updated_at_utc = ? WHERE id = ? AND user_id = ?",
            (utc_now_iso(), conversation_id, user_id),
        )
        conn.commit()
    finally:
        conn.close()


def ai_edit_user_message(message_id: int, user_id: int, new_content: str) -> dict[str, Any] | None:
    new_content = (new_content or "").strip()
    if not new_content:
        return None
    conn = get_conn()
    try:
        cur = conn.cursor()
        cur.execute(
            """
            SELECT m.conversation_id
            FROM ai_messages m
            JOIN ai_conversations c ON c.id = m.conversation_id
            WHERE m.id = ? AND m.role = 'user' AND c.user_id = ?
            """,
            (int(message_id), user_id),
        )
        row = cur.fetchone()
        if not row:
            return None
        conv_id = int(row[0])
        edited_at = utc_now_iso()
        cur.execute("UPDATE ai_messages SET content = ?, edited_at_utc = ? WHERE id = ?", (new_content, edited_at, int(message_id)))
        # Editing a user message invalidates all messages after it.
        cur.execute("DELETE FROM ai_messages WHERE conversation_id = ? AND id > ?", (conv_id, int(message_id)))
        cur.execute("UPDATE ai_conversations SET updated_at_utc = ? WHERE id = ? AND user_id = ?", (edited_at, conv_id, user_id))
        conn.commit()
        return {"conversation_id": conv_id, "message_id": int(message_id), "edited_at_utc": edited_at}
    finally:
        conn.close()


def ai_update_conversation(
    conversation_id: int,
    user_id: int,
    *,
    title: str | None = None,
    mode: str | None = None,
    system_prompt: str | None = None,
    is_pinned: bool | None = None,
    is_bookmarked: bool | None = None,
) -> dict[str, Any] | None:
    if not _ai_conv_belongs_to_user(conversation_id, user_id):
        return None
    fields = []
    values: list[Any] = []
    if title is not None:
        t = title.strip()
        if t:
            fields.append("title = ?")
            values.append(t)
    if mode is not None:
        m = (mode or "").strip().lower()
        if m in CHAT_MODES:
            fields.append("mode = ?")
            values.append(m)
    if system_prompt is not None:
        fields.append("system_prompt = ?")
        values.append((system_prompt or "").strip())
    if is_pinned is not None:
        fields.append("is_pinned = ?")
        values.append(1 if is_pinned else 0)
    if is_bookmarked is not None:
        fields.append("is_bookmarked = ?")
        values.append(1 if is_bookmarked else 0)
    if not fields:
        return ai_conversation_brief(conversation_id, user_id)
    now = utc_now_iso()
    fields.append("updated_at_utc = ?")
    values.append(now)
    values.extend([conversation_id, user_id])
    conn = get_conn()
    try:
        cur = conn.cursor()
        cur.execute(f"UPDATE ai_conversations SET {', '.join(fields)} WHERE id = ? AND user_id = ?", tuple(values))
        conn.commit()
    finally:
        conn.close()
    return ai_conversation_brief(conversation_id, user_id)


def ai_delete_conversation(conversation_id: int, user_id: int) -> bool:
    if not _ai_conv_belongs_to_user(conversation_id, user_id):
        return False
    conn = get_conn()
    try:
        cur = conn.cursor()
        cur.execute("DELETE FROM ai_messages WHERE conversation_id = ?", (conversation_id,))
        cur.execute("DELETE FROM ai_conversations WHERE id = ? AND user_id = ?", (conversation_id, user_id))
        conn.commit()
        return True
    finally:
        conn.close()


def ai_build_messages_for_model(conversation_id: int, user_id: int) -> tuple[str, list[dict[str, str]]]:
    """
    Returns (mode, messages) ready for OpenAI Chat Completions.
    Trims long history to keep latency/cost reasonable.
    """
    conn = get_conn()
    try:
        cur = conn.cursor()
        cur.execute(
            "SELECT mode, COALESCE(system_prompt, ''), COALESCE(summary, '') FROM ai_conversations WHERE id = ? AND user_id = ?",
            (conversation_id, user_id),
        )
        row = cur.fetchone()
        if not row:
            raise ValueError("Conversation not found")
        mode = (row[0] or "normal").strip().lower()
        if mode not in CHAT_MODES:
            mode = "normal"
        custom_system = (row[1] or "").strip()
        summary = (row[2] or "").strip()
        cur.execute(
            """
            SELECT role, content
            FROM ai_messages
            WHERE conversation_id = ?
            ORDER BY id DESC
            LIMIT ?
            """,
            (conversation_id, int(AI_HISTORY_LIMIT * 2)),
        )
        rows = cur.fetchall()
        rows.reverse()
    finally:
        conn.close()

    system_text = custom_system or CHAT_MODES[mode]
    memories_txt = ""
    if AI_MEMORY_LIMIT > 0:
        mems = list_user_memories(user_id, limit=AI_MEMORY_LIMIT)
        if mems:
            bullets = "\n".join([f"- {m['text']}" for m in reversed(mems)])
            memories_txt = "User memory (facts that may help):\n" + bullets

    system_parts = [system_text]
    if memories_txt:
        system_parts.append(memories_txt)
    if summary:
        system_parts.append("Conversation summary:\n" + summary)

    msgs = [{"role": "system", "content": "\n\n".join(system_parts)}]
    for r in rows:
        role = (r[0] or "").strip().lower()
        if role in {"user", "assistant"}:
            msgs.append({"role": role, "content": r[1] or ""})
    return mode, msgs


def normalize_tz_name(tz_name: str | None) -> str | None:
    tz_name = (tz_name or "").strip()
    if not tz_name:
        return None
    if ZoneInfo is None:
        return None
    try:
        ZoneInfo(tz_name)
        return tz_name
    except Exception:
        return None


def usage_day_iso(tz_name: str | None) -> str:
    tz = normalize_tz_name(tz_name)
    if tz:
        return datetime.now(ZoneInfo(tz)).date().isoformat()
    # Fallback: server day (better than failing hard).
    return date.today().isoformat()


def next_reset_iso(tz_name: str | None) -> str | None:
    tz = normalize_tz_name(tz_name)
    if tz:
        now = datetime.now(ZoneInfo(tz))
        tomorrow = now.date() + timedelta(days=1)
        reset_at = datetime.combine(tomorrow, time.min, tzinfo=ZoneInfo(tz))
        return reset_at.isoformat()
    # Fallback: next UTC midnight (keeps client lock working even without tz support).
    now_utc = datetime.now(timezone.utc)
    tomorrow_utc = now_utc.date() + timedelta(days=1)
    reset_at_utc = datetime.combine(tomorrow_utc, time.min, tzinfo=timezone.utc)
    return reset_at_utc.isoformat()


def get_user_plan(user_id: int) -> str:
    conn = get_conn()
    try:
        cur = conn.cursor()
        cur.execute("SELECT plan FROM users WHERE id = ?", (user_id,))
        row = cur.fetchone()
        plan = (row[0] if row else None) or FREE_PLAN
        return str(plan).strip().lower() or FREE_PLAN
    finally:
        conn.close()


def get_user_tz(user_id: int) -> str | None:
    conn = get_conn()
    try:
        cur = conn.cursor()
        cur.execute("SELECT tz FROM users WHERE id = ?", (user_id,))
        row = cur.fetchone()
        return normalize_tz_name(row[0] if row else None)
    finally:
        conn.close()


def set_user_tz_if_empty(user_id: int, tz_name: str) -> None:
    tz = normalize_tz_name(tz_name)
    if not tz:
        return
    conn = get_conn()
    try:
        cur = conn.cursor()
        cur.execute(
            "UPDATE users SET tz = ? WHERE id = ? AND (tz IS NULL OR TRIM(tz) = '')",
            (tz, user_id),
        )
        conn.commit()
    finally:
        conn.close()


def get_daily_message_count(user_id: int, day: str) -> int:
    conn = get_conn()
    try:
        cur = conn.cursor()
        cur.execute(
            "SELECT message_count FROM daily_message_usage WHERE user_id = ? AND day = ?",
            (user_id, day),
        )
        row = cur.fetchone()
        return int(row[0]) if row else 0
    finally:
        conn.close()


def reserve_daily_message(user_id: int, day: str) -> int:
    """
    Reserves (increments) one message slot for (user_id, day) transactionally.
    Returns the new message_count.
    """
    conn = get_conn()
    try:
        conn.execute("BEGIN IMMEDIATE")
        cur = conn.cursor()
        cur.execute(
            "SELECT message_count FROM daily_message_usage WHERE user_id = ? AND day = ?",
            (user_id, day),
        )
        row = cur.fetchone()
        current_count = int(row[0]) if row else 0
        new_count = current_count + 1
        if row:
            cur.execute(
                "UPDATE daily_message_usage SET message_count = ? WHERE user_id = ? AND day = ?",
                (new_count, user_id, day),
            )
        else:
            cur.execute(
                "INSERT INTO daily_message_usage (user_id, day, message_count) VALUES (?, ?, ?)",
                (user_id, day, new_count),
            )
        conn.commit()
        return new_count
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()


def release_daily_message(user_id: int, day: str) -> None:
    """Best-effort rollback: decrement usage for (user_id, day) if possible."""
    conn = get_conn()
    try:
        conn.execute("BEGIN IMMEDIATE")
        cur = conn.cursor()
        cur.execute(
            "SELECT message_count FROM daily_message_usage WHERE user_id = ? AND day = ?",
            (user_id, day),
        )
        row = cur.fetchone()
        if not row:
            conn.rollback()
            return
        current_count = int(row[0])
        new_count = max(0, current_count - 1)
        cur.execute(
            "UPDATE daily_message_usage SET message_count = ? WHERE user_id = ? AND day = ?",
            (new_count, user_id, day),
        )
        conn.commit()
    except Exception:
        conn.rollback()
    finally:
        conn.close()


def is_admin_request() -> bool:
    admin_token = (os.getenv("ADMIN_TOKEN") or "").strip()
    if not admin_token:
        return False
    provided = (request.headers.get("X-Admin-Token") or "").strip()
    if provided and provided == admin_token:
        return True
    # Convenience for browser usage; avoid using this in public deployments.
    provided_qs = (request.args.get("token", "") or "").strip()
    return bool(provided_qs) and provided_qs == admin_token


def admin_required() -> tuple[bool, Response | None]:
    admin_token = (os.getenv("ADMIN_TOKEN") or "").strip()
    if not admin_token:
        return False, jsonify({"error": "Admin endpoints are disabled."})
    if not is_admin_request():
        return False, jsonify({"error": "Forbidden."})
    return True, None


def canned_about_app_reply(user_message: str) -> str | None:
    """
    Deterministic answers about this app/company.
    Requirement: regardless of phrasing, if the intent matches, always return the configured answer.
    """

    def norm(s: str) -> str:
        s = (s or "").strip().lower()
        if not s:
            return ""
        # Keep letters/numbers from all scripts; strip punctuation to make matching robust.
        s = re.sub(r"[^\w\s]+", " ", s, flags=re.UNICODE)
        s = re.sub(r"\s+", " ", s, flags=re.UNICODE).strip()
        return s

    def has_any(s: str, toks: list[str]) -> bool:
        return any(t and (t in s) for t in toks)

    def mentions_name(s: str) -> bool:
        variants = [
            ("erik", "petrosyan"),
            ("էրիկ", "պետրոսյան"),
            ("երիկ", "պետրոսյան"),
            ("эрик", "петросян"),
        ]
        return any(a in s and b in s for a, b in variants)

    def detect_lang(raw: str) -> str:
        r = raw or ""
        if re.search(r"[\u0530-\u058F]", r):  # Armenian
            return "hy"
        if re.search(r"[\u0600-\u06FF]", r):  # Arabic/Persian
            return "fa"
        if re.search(r"[\u0400-\u04FF]", r):  # Cyrillic (ru/uk)
            # crude: Ukrainian unique letters
            if re.search(r"[іїєґІЇЄҐ]", r):
                return "uk"
            return "ru"
        # fallback by keywords
        low = (r or "").lower()
        if re.search(r"(?:^|\W)(wer|wem|erstellt|erschaffen|entwickelt|autor)(?:$|\W)", low):
            return "de"
        return "en"

    def questionish(raw: str, s: str) -> bool:
        if "?" in (raw or ""):
            return True
        return has_any(
            s,
            [
                "who",
                "who is",
                "tell me",
                "about",
                "what is",
                "kto",
                "кто",
                "хто",
                "ov",
                "ով",
                "ինչ",
                "ասա",
                "պատմ",
            ],
        )

    raw = (user_message or "").strip()
    s = norm(raw)
    if not s:
        return None

    lang = detect_lang(raw)

    def reply_company_owner(lang_code: str) -> str:
        # Keep it short, and never mention OpenAI/ChatGPT per product requirement.
        if lang_code == "hy":
            return "YAN Company։ Տեր՝ Erik Petrosyan։"
        if lang_code == "ru":
            return "YAN Company. Владелец: Erik Petrosyan."
        if lang_code == "uk":
            return "YAN Company. Власник: Erik Petrosyan."
        if lang_code == "de":
            return "YAN Company. Inhaber: Erik Petrosyan."
        if lang_code == "fa":
            return "YAN Company. مالک: Erik Petrosyan."
        return "YAN Company. Owner: Erik Petrosyan."

    # Intent: "who created you / who made yan"
    creator_tokens = [
        # English
        "create",
        "created",
        "creator",
        "made",
        "make",
        "built",
        "build",
        "develop",
        "developed",
        # Armenian stems
        "ստեղծ",
        "սարք",
        "հեղինակ",
        # Armenian (latin translit stems)
        "stex",
        "stexc",
        "stexcel",
        "stexcvel",
        "sarq",
        "heghinak",
        "hexinak",
        # Russian/Ukrainian stems
        "созд",
        "сдел",
        "разработ",
        "створ",
        "зроб",
        # Persian stems
        "ساخت",
        "سازنده",
        "توسعه",
    ]
    you_tokens = [
        "you",
        "u",
        "yan",
        # Armenian
        "քեզ",
        "ձեզ",
        "դու",
        # Armenian translit
        "qez",
        "dzez",
        "du",
        # Russian/Ukrainian
        "теб",
        "вас",
        "ты",
        "ви",
        "тебе",
        "тоб",
        "тобі",
        # Persian
        "تو",
        "شما",
    ]
    q_tokens = ["who", "кто", "хто", "ov", "ով"]

    is_creator_question = (has_any(s, creator_tokens) and has_any(s, you_tokens) and (has_any(s, q_tokens) or "?" in raw or len(s.split()) <= 7))
    if is_creator_question:
        return reply_company_owner(lang)

    # Intent: "are you ChatGPT/OpenAI/GPT" / "who are you (model/provider)".
    provider_tokens = [
        "openai",
        "chatgpt",
        "gpt",
        "gpt4",
        "gpt-4",
        "gpt 4",
        "assistant",
        "model",
        "llm",
        # Russian/Ukrainian
        "чатгпт",
        "чат gpt",
        "опенай",
        "модель",
        # Armenian
        "չաթգպտ",
        "օփենայ",
    ]
    identity_q_tokens = [
        "are you",
        "what are you",
        "who are you",
        "is this",
        "is it",
        "you are",
        # Armenian
        "դու",
        "դուք",
        "ո՞վ ես",
        "ինչ ես",
        # Russian/Ukrainian
        "ты",
        "вы",
        "кто ты",
        "что ты",
        "ти",
        "хто ти",
        "що ти",
        # Persian
        "تو",
        "شما",
    ]
    if (has_any(s, provider_tokens) and (questionish(raw, s) or has_any(s, identity_q_tokens)) and has_any(s, you_tokens)):
        return reply_company_owner(lang)

    # Intent: "who is Erik Petrosyan" (or equivalent)
    if mentions_name(s) and questionish(raw, s):
        return reply_company_owner(lang)

    return None


def get_user_daily_limit_override(user_id: int) -> int | None:
    """Per-user override for FREE plan daily message limit. None means use global FREE_DAILY_MESSAGE_LIMIT."""
    conn = get_conn()
    try:
        cur = conn.cursor()
        cur.execute("SELECT daily_limit_override FROM users WHERE id = ?", (int(user_id),))
        row = cur.fetchone()
        if not row:
            return None
        v = row[0]
        if v is None:
            return None
        try:
            n = int(v)
        except Exception:
            return None
        return n if n >= 0 else None
    finally:
        conn.close()


def set_user_daily_limit_override(user_id: int, limit: int | None) -> None:
    conn = get_conn()
    try:
        cur = conn.cursor()
        if limit is None:
            cur.execute("UPDATE users SET daily_limit_override = NULL WHERE id = ?", (int(user_id),))
        else:
            cur.execute("UPDATE users SET daily_limit_override = ? WHERE id = ?", (int(limit), int(user_id)))
        conn.commit()
    finally:
        conn.close()


def effective_free_daily_limit_for_user(user_id: int) -> int:
    """Effective free daily limit for this user (override or global)."""
    override = get_user_daily_limit_override(int(user_id))
    if override is None:
        return int(FREE_DAILY_MESSAGE_LIMIT)
    return int(max(0, override))


def _firebase_uid_for_user_id(user_id: int) -> str:
    conn = get_conn()
    try:
        cur = conn.cursor()
        cur.execute(
            "SELECT provider_user_id FROM identities WHERE user_id = ? AND provider = ? LIMIT 1",
            (int(user_id), "firebase"),
        )
        row = cur.fetchone()
        return str(row[0] or "").strip() if row else ""
    finally:
        conn.close()


def _billing_clear_subscription_local(user_id: int) -> None:
    conn = get_conn()
    try:
        cur = conn.cursor()
        cur.execute("DELETE FROM billing_subscriptions WHERE user_id = ?", (int(user_id),))
        conn.commit()
    finally:
        conn.close()


def _billing_force_revoke_firestore_entitlements(firebase_uid: str) -> None:
    uid = (firebase_uid or "").strip()
    if not uid:
        return
    doc_ref = _billing_firestore_doc(uid)
    if doc_ref is None:
        return
    try:
        doc_ref.set(
            {
                "uid": uid,
                "plan_type": FREE_PLAN,
                "subscription_status": "",
                "stripe_customer_id": "",
                "stripe_subscription_id": "",
                "current_period_end_utc": "",
                "cancel_at_period_end": False,
                "updated_at_utc": utc_now_iso(),
            },
            merge=True,
        )
    except Exception:
        return


def allowed_image(filename: str) -> bool:
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    return ext in ALLOWED_IMAGE_EXTENSIONS


def save_profile_photo(file_storage) -> str | None:
    if not file_storage or not file_storage.filename:
        return None
    filename = secure_filename(file_storage.filename)
    if not filename or not allowed_image(filename):
        raise ValueError("Unsupported image format. Use PNG, JPG, JPEG, GIF, or WEBP.")
    ext = filename.rsplit(".", 1)[-1].lower()
    generated = f"user_{uuid4().hex}.{ext}"
    abs_path = os.path.join(UPLOAD_FOLDER, generated)
    file_storage.save(abs_path)
    return f"uploads/{generated}"


def get_user_profile(user_id: int) -> dict:
    conn = get_conn()
    cur = conn.cursor()
    cur.execute(
        "SELECT id, username, email, address, age, photo_path FROM users WHERE id = ?",
        (user_id,),
    )
    row = cur.fetchone()
    conn.close()
    if not row:
        return {}
    return {
        "id": row[0],
        "username": row[1] or "",
        "email": row[2] or "",
        "address": row[3] or "",
        "age": row[4] if row[4] is not None else "",
        "photo_path": row[5] or "",
    }


def username_from_profile(display_name, email):
    base = (display_name or email or "user").lower()
    base = re.sub(r"[^a-z0-9_]+", "_", base).strip("_")
    if len(base) < 3:
        base = "user"
    return base[:24]


def make_unique_username(base_username):
    conn = get_conn()
    cur = conn.cursor()
    candidate = base_username
    suffix = 1
    while True:
        cur.execute("SELECT 1 FROM users WHERE username = ?", (candidate,))
        exists = cur.fetchone()
        if not exists:
            conn.close()
            return candidate
        suffix += 1
        candidate = f"{base_username}_{suffix}"


def login_with_identity(provider, provider_user_id, email, display_name):
    conn = get_conn()
    cur = conn.cursor()
    cur.execute(
        "SELECT user_id FROM identities WHERE provider = ? AND provider_user_id = ?",
        (provider, provider_user_id),
    )
    row = cur.fetchone()

    if row:
        user_id = row[0]
        cur.execute("SELECT username FROM users WHERE id = ?", (user_id,))
        user_row = cur.fetchone()
        conn.close()
        if not user_row:
            return False
        session["user_id"] = user_id
        session["username"] = user_row[0]
        conversations.setdefault(user_id, {})
        return True

    base_username = username_from_profile(display_name, email)
    unique_username = make_unique_username(base_username)
    cur.execute(
        "INSERT INTO users (username, password_hash) VALUES (?, ?)",
        (unique_username, "oauth-login"),
    )
    user_id = cur.lastrowid
    cur.execute(
        "INSERT INTO identities (user_id, provider, provider_user_id, email) VALUES (?, ?, ?, ?)",
        (user_id, provider, provider_user_id, email),
    )
    conn.commit()
    conn.close()

    session["user_id"] = user_id
    session["username"] = unique_username
    conversations[user_id] = {}
    return True


def get_user_chats(user_id):
    user_chats = conversations.setdefault(user_id, {})
    if isinstance(user_chats, list):
        user_chats = {"default": user_chats}
        conversations[user_id] = user_chats
    return user_chats


init_db()
configure_oauth()


@app.after_request
def _no_store_html(resp: Response):
    # Make UI iterations visible without restarting the server.
    try:
        if (resp.mimetype or "").startswith("text/html"):
            resp.headers["Cache-Control"] = "no-store"
            # Firebase Auth popup flow (Google sign-in) works best when the opener
            # page allows popups to keep a link to the opener.
            resp.headers.setdefault("Cross-Origin-Opener-Policy", "same-origin-allow-popups")
            # Ensure we don't accidentally opt into COEP which can break third-party auth SDKs.
            resp.headers.setdefault("Cross-Origin-Embedder-Policy", "unsafe-none")
    except Exception:
        pass
    return resp


def login_required(func):
    @wraps(func)
    def wrapper(*args, **kwargs):
        if "user_id" not in session:
            if request.path.startswith("/chat") or request.path.startswith("/reset"):
                return jsonify({"error": "Please login first."}), 401
            return redirect(url_for("auth_page"))
        return func(*args, **kwargs)

    return wrapper


@app.route("/auth")
def auth_page():
    if "user_id" in session:
        return redirect(url_for("app_home"))
    fb_err = _firebase_init_error()
    return render_template(
        "auth.html",
        error=request.args.get("error", ""),
        message=request.args.get("message", ""),
        google_enabled="google" in oauth_clients,
        microsoft_enabled="microsoft" in oauth_clients,
        firebase_config=_firebase_web_config(),
        firebase_error=fb_err,
    )


@app.route("/oauth/<provider>")
def oauth_login(provider):
    remote = oauth_clients.get(provider)
    if remote is None:
        return redirect(url_for("auth_page", error=f"{provider.capitalize()} login is not configured yet."))
    redirect_uri = url_for("oauth_callback", provider=provider, _external=True)
    return remote.authorize_redirect(redirect_uri)


@app.route("/oauth/<provider>/callback")
def oauth_callback(provider):
    remote = oauth_clients.get(provider)
    if remote is None:
        return redirect(url_for("auth_page", error=f"{provider.capitalize()} login is not configured yet."))

    try:
        token = remote.authorize_access_token()
        userinfo = None

        try:
            userinfo = remote.parse_id_token(token)
        except Exception:
            userinfo = token.get("userinfo")

        if provider == "microsoft" and not userinfo:
            resp = remote.get("https://graph.microsoft.com/oidc/userinfo")
            userinfo = resp.json()
        elif provider == "google" and not userinfo:
            resp = remote.get("userinfo")
            userinfo = resp.json()

        if not userinfo:
            return redirect(url_for("auth_page", error="Could not fetch user profile from provider."))

        if provider == "microsoft":
            provider_user_id = userinfo.get("sub") or userinfo.get("oid")
            email = userinfo.get("email") or userinfo.get("preferred_username")
            display_name = userinfo.get("name") or userinfo.get("preferred_username")
        else:
            provider_user_id = userinfo.get("sub")
            email = userinfo.get("email")
            display_name = userinfo.get("name")

        if not provider_user_id:
            return redirect(url_for("auth_page", error="Provider did not return a valid user id."))

        if not login_with_identity(provider, provider_user_id, email, display_name):
            return redirect(url_for("auth_page", error="Login failed while creating user profile."))

        session.permanent = True
        return redirect(url_for("app_home"))
    except Exception as exc:
        return redirect(url_for("auth_page", error=f"OAuth failed: {str(exc)}"))


@app.route("/firebase/session", methods=["POST"])
def firebase_session_login():
    # Accept Firebase ID token from the browser; verify with Firebase Admin; map to local user.
    if not request.is_json:
        return jsonify({"ok": False, "error": "Request must be JSON"}), 400

    fb_err = _firebase_init_error()
    if fb_err:
        return jsonify({"ok": False, "error": fb_err}), 500

    payload = request.get_json(silent=True) or {}
    id_token = str(payload.get("idToken") or "").strip()
    nickname = str(payload.get("nickname") or "").strip()
    if not id_token:
        return jsonify({"ok": False, "error": "Missing idToken"}), 400
    if nickname and (not re.fullmatch(r"[A-Za-z0-9_]{3,24}", nickname)):
        return jsonify({"ok": False, "error": "Nickname must be 3-24 characters: letters, numbers, underscore."}), 400

    try:
        # Browsers/phones and servers can be off by a second; allow small leeway.
        skew = 5
        try:
            skew = int(str(os.getenv("FIREBASE_CLOCK_SKEW_SECONDS", "5") or "5").strip() or "5")
        except Exception:
            skew = 5
        decoded = fb_auth.verify_id_token(id_token, check_revoked=True, clock_skew_seconds=max(0, min(skew, 60)))  # type: ignore[union-attr]
    except Exception as exc:
        return jsonify({"ok": False, "error": f"Invalid token: {str(exc)}"}), 401

    uid = str(decoded.get("uid") or decoded.get("user_id") or "").strip()
    email = str(decoded.get("email") or "").strip().lower()
    display_name = str(decoded.get("name") or decoded.get("displayName") or "").strip()
    firebase_claim = decoded.get("firebase") if isinstance(decoded.get("firebase"), dict) else {}
    sign_in_provider = str((firebase_claim or {}).get("sign_in_provider") or "").strip().lower()
    if not uid:
        return jsonify({"ok": False, "error": "Token missing uid"}), 401
    # Don't over-validate emails here. Firebase already validates emails on signup,
    # and legitimate domains can be dotless (e.g. internal/local dev).
    def _looks_like_email(s: str) -> bool:
        s = (s or "").strip()
        return bool(s) and ("@" in s) and (not re.search(r"\\s", s))

    if not _looks_like_email(email):
        # Some auth methods (anonymous / phone) legitimately have no email.
        # We still try to fetch the user record because it may contain an email.
        try:
            u = fb_auth.get_user(uid)  # type: ignore[union-attr]
            email = str(getattr(u, "email", "") or "").strip().lower()
            if not display_name:
                display_name = str(getattr(u, "display_name", "") or "").strip()
        except Exception:
            # If lookup fails, we can still proceed with UID mapping.
            email = ""

    ip = client_ip()
    now_iso = utc_now_iso()

    conn = get_conn()
    try:
        cur = conn.cursor()
        user_id: int = 0
        username = ""
        if _looks_like_email(email):
            cur.execute("SELECT id, username FROM users WHERE email = ? COLLATE NOCASE LIMIT 1", (email,))
            r = cur.fetchone()
            if r:
                user_id = int(r[0])
                username = r[1] or ""
        else:
            # No email: map by Firebase UID via identities table.
            cur.execute(
                "SELECT user_id FROM identities WHERE provider = ? AND provider_user_id = ? LIMIT 1",
                ("firebase", uid),
            )
            id_row = cur.fetchone()
            if id_row:
                user_id = int(id_row[0])
                cur.execute("SELECT username, COALESCE(email, '') FROM users WHERE id = ? LIMIT 1", (user_id,))
                urow = cur.fetchone() or ("", "")
                username = str(urow[0] or "")
                if not _looks_like_email(email):
                    email = str(urow[1] or "").strip().lower()

        if not user_id:
            if nickname:
                base = nickname.lower()
            else:
                base = username_from_profile(display_name, email or f"guest_{uid[:6]}")
            username = make_unique_username(base)
            cur.execute(
                "INSERT INTO users (username, password_hash, email, created_at_utc, password_changed_at_utc, failed_login_count, locked_until_utc) "
                "VALUES (?, ?, ?, ?, ?, 0, '')",
                (username, "firebase", email, now_iso, now_iso),
            )
            conn.commit()
            user_id = int(cur.lastrowid)

        # Store firebase uid in identities table to support UID mapping/audits.
        try:
            cur.execute(
                "INSERT OR IGNORE INTO identities (user_id, provider, provider_user_id, email) VALUES (?, ?, ?, ?)",
                (user_id, "firebase", uid, email),
            )
        except Exception:
            pass

        cur.execute(
            "UPDATE users SET last_login_at_utc = ?, last_login_ip = ?, failed_login_count = 0, locked_until_utc = '' WHERE id = ?",
            (now_iso, ip, user_id),
        )
        conn.commit()
    finally:
        conn.close()

    session["user_id"] = user_id
    session["username"] = username
    session["firebase_uid"] = uid
    session["email"] = email
    session["sign_in_provider"] = sign_in_provider
    session["is_guest"] = (sign_in_provider == "anonymous")
    # Prefer a human-friendly label for UI (Firebase displayName / signup nickname),
    # while keeping `username` as the stable local identifier.
    ui_name = nickname or display_name
    if not ui_name:
        # Avoid showing the full email in UI if no nickname exists yet.
        ui_name = (email.split("@", 1)[0] if _looks_like_email(email) else username)
    session["display_name"] = ui_name
    session.permanent = True
    conversations.setdefault(user_id, {})

    # Best-effort: seed Firestore billing doc for this Firebase user (plan/usage/subscription).
    try:
        _billing_firestore_ensure_user(uid, email, user_id=int(user_id), plan_hint=get_user_plan(int(user_id)))
    except Exception:
        pass

    return jsonify({"ok": True, "user_id": user_id, "username": username})


@app.route("/register", methods=["POST"])
def register():
    # Deprecated: auth is handled by Firebase.
    return redirect(url_for("auth_page", error="Registration is handled by Firebase. Please sign in with Firebase."))


@app.route("/login", methods=["POST"])
def login():
    # Deprecated: auth is handled by Firebase.
    return redirect(url_for("auth_page", error="Login is handled by Firebase. Please sign in with Firebase."))


@app.route("/logout", methods=["POST"])
def logout():
    require_csrf()
    user_id = session.get("user_id")
    if user_id in conversations:
        conversations.pop(user_id, None)
    session.clear()
    return redirect(url_for("auth_page"))


@app.route("/email-login/request", methods=["POST"])
def email_login_request():
    return redirect(url_for("auth_page", error="Email OTP login was replaced by Firebase."))


@app.route("/email-login/verify", methods=["POST"])
def email_login_verify():
    return redirect(url_for("auth_page", error="Email OTP login was replaced by Firebase."))


@app.route("/password-reset", methods=["GET", "POST"])
def password_reset_request():
    return redirect(url_for("auth_page", error="Password reset is handled by Firebase."))


@app.route("/password-reset/<token>", methods=["GET", "POST"])
def password_reset_set(token: str):
    return redirect(url_for("auth_page", error="Password reset is handled by Firebase."))


@app.route("/")
def landing():
    if "user_id" in session:
        return redirect(url_for("app_home"))
    # Public landing page so crawlers (and users) get real content at `/`.
    # The actual app UI still requires auth at `/ai`.
    return render_template("landing.html")


@app.route("/app")
@login_required
def app_home():
    # Single source of truth for the main AI chat UI.
    resp = redirect(url_for("ai_home"))
    # Avoid browsers/proxies caching the redirect and/or an older rendered page.
    resp.headers["Cache-Control"] = "no-store, max-age=0"
    resp.headers["Pragma"] = "no-cache"
    return resp


@app.route("/ai")
@login_required
def ai_home():
    # Prefer the React/Tailwind SPA build if present.
    # You can force behavior via env:
    # - YAN_USE_SPA=1  => always try SPA first
    # - YAN_USE_SPA=0  => always use server-rendered template UI
    use_spa_raw = str(os.getenv("YAN_USE_SPA", "") or "").strip().lower()
    use_spa_forced = None
    if use_spa_raw in ("1", "true", "yes", "on"):
        use_spa_forced = True
    elif use_spa_raw in ("0", "false", "no", "off"):
        use_spa_forced = False

    manifest_path = Path(app.static_folder or "static") / "spa" / ".vite" / "manifest.json"
    use_spa = manifest_path.exists() if use_spa_forced is None else use_spa_forced
    if use_spa and manifest_path.exists():
        try:
            manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
            entry = manifest.get("index.html") or {}
            js_file = str(entry.get("file") or "").strip()
            css_files = entry.get("css") or []
            css_files = [str(x) for x in css_files if str(x).strip()]
            if js_file:
                return render_template(
                    "ai_spa.html",
                    username=session.get("username", "User"),
                    display_name=session.get("display_name", ""),
                    vite_js=f"spa/{js_file}".replace("\\", "/"),
                    vite_css=[f"spa/{c}".replace("\\", "/") for c in css_files],
                )
        except Exception:
            pass

    return render_template(
        "ai.html",
        username=session.get("username", "User"),
        display_name=session.get("display_name", ""),
        spa_missing=not manifest_path.exists(),
    )


@app.route("/settings")
@login_required
def settings_page():
    if session.get("is_guest"):
        return render_template("guest_register_required.html", title="Settings - yan")
    user_id = int(session["user_id"])
    profile = get_user_profile(user_id)
    billing = billing_status_for_user(user_id, requested_tz=None)
    return render_template("settings.html", profile=profile, billing=billing)


@app.route("/upgrade")
@login_required
def upgrade_page():
    user_id = int(session["user_id"])
    billing = billing_status_for_user(user_id, requested_tz=None)
    return render_template(
        "upgrade.html",
        stripe_enabled=_stripe_enabled(),
        current_plan=str((billing or {}).get("plan") or FREE_PLAN),
        free_daily_limit=int(FREE_DAILY_MESSAGE_LIMIT),
        prices=_stripe_price_ids(),
        billing=billing,
    )


@app.route("/api/billing/status", methods=["GET"])
@login_required
def api_billing_status():
    user_id = int(session["user_id"])
    tz = str(request.args.get("tz") or "").strip() or None
    return jsonify({"billing": billing_status_for_user(user_id, requested_tz=tz)})


@app.route("/api/billing/checkout", methods=["POST"])
@login_required
def api_billing_checkout():
    if not request.is_json:
        return jsonify({"error": "Request must be JSON."}), 400
    if session.get("is_guest"):
        return jsonify({"error": "Guest accounts can't upgrade. Create an account to continue."}), 403
    if not _stripe_enabled():
        return jsonify({"error": "Stripe billing isn't configured on the server."}), 500

    payload = request.get_json(silent=True) or {}
    plan = str(payload.get("plan") or "pro").strip().lower() or "pro"
    interval = str(payload.get("interval") or "month").strip().lower() or "month"
    if interval in {"monthly", "month"}:
        interval = "month"
    elif interval in {"yearly", "year"}:
        interval = "year"
    else:
        return jsonify({"error": "Unsupported interval."}), 400

    prices = _stripe_price_ids()
    price_id = str(((prices.get(plan) or {}).get(interval) or "")).strip()
    if not price_id:
        return jsonify({"error": "This plan is not configured yet (missing Stripe price ID)."}), 400

    user_id = int(session["user_id"])
    firebase_uid = str(session.get("firebase_uid") or "").strip()
    email = str(session.get("email") or "").strip().lower()
    if not firebase_uid:
        return jsonify({"error": "Missing Firebase UID on session. Log out and sign in again."}), 400

    _stripe_set_api_key()

    # Get or create Stripe customer.
    customer_id = _user_get_stripe_customer_id(user_id)
    if not customer_id:
        # Try Firestore first if present.
        doc = _billing_firestore_ensure_user(firebase_uid, email, user_id=user_id, plan_hint=get_user_plan(user_id))
        customer_id = str((doc or {}).get("stripe_customer_id") or "").strip()
    if not customer_id:
        c = stripe.Customer.create(  # type: ignore[union-attr]
            email=email or None,
            metadata={"firebase_uid": firebase_uid, "local_user_id": str(user_id)},
        )
        customer_id = str(c.get("id") or "").strip()
        if customer_id:
            _user_set_stripe_customer_id(user_id, customer_id)
            try:
                doc_ref = _billing_firestore_doc(firebase_uid)
                if doc_ref is not None:
                    doc_ref.set({"stripe_customer_id": customer_id, "updated_at_utc": _billing_firestore_now_iso()}, merge=True)
            except Exception:
                pass

    if not customer_id:
        return jsonify({"error": "Failed to create Stripe customer."}), 500

    # Best-effort: ensure customer metadata includes a stable mapping back to this user.
    try:
        stripe.Customer.modify(  # type: ignore[union-attr]
            customer_id,
            metadata={"firebase_uid": firebase_uid, "local_user_id": str(user_id)},
        )
    except Exception:
        pass

    base = _best_public_base_url().rstrip("/") or request.url_root.rstrip("/")
    success_url = base + "/upgrade?success=1"
    cancel_url = base + "/upgrade?canceled=1"

    s = stripe.checkout.Session.create(  # type: ignore[union-attr]
        mode="subscription",
        customer=customer_id,
        line_items=[{"price": price_id, "quantity": 1}],
        allow_promotion_codes=True,
        success_url=success_url,
        cancel_url=cancel_url,
        client_reference_id=firebase_uid,
        metadata={"firebase_uid": firebase_uid, "local_user_id": str(user_id), "plan": plan},
        subscription_data={"metadata": {"firebase_uid": firebase_uid, "local_user_id": str(user_id), "plan": plan}},
    )
    url = str(s.get("url") or "").strip()
    if not url:
        return jsonify({"error": "Stripe checkout did not return a redirect URL."}), 500
    return jsonify({"url": url})


@app.route("/api/billing/portal", methods=["POST"])
@login_required
def api_billing_portal():
    if session.get("is_guest"):
        return jsonify({"error": "Guest accounts don't have billing. Create an account to continue."}), 403
    if not _stripe_enabled():
        return jsonify({"error": "Stripe billing isn't configured on the server."}), 500

    user_id = int(session["user_id"])
    firebase_uid = str(session.get("firebase_uid") or "").strip()
    email = str(session.get("email") or "").strip().lower()
    if not firebase_uid:
        return jsonify({"error": "Missing Firebase UID on session. Log out and sign in again."}), 400

    _stripe_set_api_key()
    customer_id = _user_get_stripe_customer_id(user_id)
    if not customer_id:
        doc = _billing_firestore_ensure_user(firebase_uid, email, user_id=user_id, plan_hint=get_user_plan(user_id))
        customer_id = str((doc or {}).get("stripe_customer_id") or "").strip()
    if not customer_id:
        return jsonify({"error": "No Stripe customer found for this account yet."}), 400

    base = _best_public_base_url().rstrip("/") or request.url_root.rstrip("/")
    return_url = base + "/account"

    portal = stripe.billing_portal.Session.create(  # type: ignore[union-attr]
        customer=customer_id,
        return_url=return_url,
    )
    url = str(portal.get("url") or "").strip()
    if not url:
        return jsonify({"error": "Stripe portal did not return a URL."}), 500
    return jsonify({"url": url})


@app.route("/webhooks/stripe", methods=["POST"])
def stripe_webhook():
    """
    Stripe webhook endpoint. Configure this URL in Stripe Dashboard and set STRIPE_WEBHOOK_SECRET.
    """
    if stripe is None:
        return Response("stripe not installed", status=500)
    secret = _stripe_webhook_secret()
    if not secret:
        return Response("webhook secret not configured", status=500)
    _stripe_set_api_key()

    payload = request.get_data(cache=False, as_text=False)
    sig = request.headers.get("Stripe-Signature", "")
    try:
        event = stripe.Webhook.construct_event(payload=payload, sig_header=sig, secret=secret)  # type: ignore[union-attr]
    except Exception as exc:
        return Response(f"bad signature: {str(exc)}", status=400)

    etype = str(event.get("type") or "").strip()
    obj = (event.get("data") or {}).get("object") if isinstance(event.get("data"), dict) else None
    if not isinstance(obj, dict):
        return jsonify({"ok": True})

    # Resolve user identifiers.
    md = obj.get("metadata") if isinstance(obj.get("metadata"), dict) else {}
    firebase_uid = str((md or {}).get("firebase_uid") or "").strip()
    local_user_id = str((md or {}).get("local_user_id") or (md or {}).get("local_userid") or "").strip() or None
    customer_id = str(obj.get("customer") or "").strip()
    if not firebase_uid and customer_id:
        firebase_uid = str(_stripe_get_customer_metadata(customer_id).get("firebase_uid") or "").strip()

    if etype in {"checkout.session.completed"}:
        subscription_id = str(obj.get("subscription") or "").strip()
        if subscription_id:
            _stripe_sync_subscription(subscription_id, customer_id, firebase_uid=firebase_uid, local_user_id=local_user_id)
        return jsonify({"ok": True})

    if etype in {"customer.subscription.created", "customer.subscription.updated", "customer.subscription.deleted"}:
        _stripe_sync_subscription_object(obj, customer_id, firebase_uid=firebase_uid, local_user_id=local_user_id)
        return jsonify({"ok": True})

    if etype in {"invoice.payment_failed", "invoice.payment_succeeded"}:
        subscription_id = str(obj.get("subscription") or "").strip()
        if subscription_id:
            _stripe_sync_subscription(subscription_id, customer_id, firebase_uid=firebase_uid, local_user_id=local_user_id)
        return jsonify({"ok": True})

    return jsonify({"ok": True})


@app.route("/api/users", methods=["GET"])
@login_required
def api_users():
    user_id = int(session["user_id"])
    q = str(request.args.get("q", "")).strip()
    try:
        limit = int(request.args.get("limit", "200"))
    except ValueError:
        limit = 200
    return jsonify({"users": list_users_for_viewer(user_id, q=q, limit=max(1, min(limit, 200)))})


@app.route("/api/debug", methods=["GET"])
@login_required
def api_debug():
    user_id = int(session["user_id"])
    conn = get_conn()
    try:
        cur = conn.cursor()
        cur.execute("SELECT COUNT(*) FROM users")
        total_users = int(cur.fetchone()[0] or 0)
    finally:
        conn.close()

    return jsonify(
        {
            "ok": True,
            "user": {"id": user_id, "username": session.get("username", "")},
            "db_path": DB_PATH,
            "total_users": total_users,
        }
    )


@app.route("/search", methods=["GET"])
@login_required
def search_page():
    user_id = int(session["user_id"])
    q = str(request.args.get("q", "")).strip()
    # Show suggested users even when q is empty (Instagram-like explore).
    results = list_users_for_viewer(user_id, q=q, limit=80)
    return render_template("search.html", q=q, results=results)


@app.route("/u/<username>", methods=["GET"])
@login_required
def user_profile_page(username: str):
    viewer_id = int(session["user_id"])
    target_id = _get_user_id_by_username(username)
    if not target_id:
        return redirect(url_for("search_page", q=username))
    brief = get_user_brief(int(target_id))
    if not brief:
        return redirect(url_for("search_page", q=username))
    counts = follow_counts(int(target_id))
    return render_template(
        "profile.html",
        profile=brief,
        counts=counts,
        is_self=(int(target_id) == viewer_id),
        is_following=is_following(viewer_id, int(target_id)),
    )


@app.route("/api/profile/<username>", methods=["GET"])
@login_required
def api_profile(username: str):
    viewer_id = int(session["user_id"])
    target_id = _get_user_id_by_username(username)
    if not target_id:
        return jsonify({"error": "User not found"}), 404
    brief = get_user_brief(int(target_id))
    if not brief:
        return jsonify({"error": "User not found"}), 404
    counts = follow_counts(int(target_id))
    return jsonify(
        {
            "user": brief,
            "counts": counts,
            "viewer": {"is_self": int(target_id) == viewer_id, "is_following": is_following(viewer_id, int(target_id))},
        }
    )


@app.route("/api/follow", methods=["POST"])
@login_required
def api_follow():
    viewer_id = int(session["user_id"])
    payload = request.get_json(silent=True) or {}
    target_id = int(payload.get("user_id") or 0)
    if not target_id:
        target_id = int(_get_user_id_by_username(str(payload.get("username") or "").strip()) or 0)
    if not target_id:
        return jsonify({"error": "User not found"}), 404
    if target_id == viewer_id:
        return jsonify({"error": "Cannot follow yourself"}), 400
    set_follow(viewer_id, target_id, True)
    return jsonify({"ok": True, "is_following": True, "counts": follow_counts(target_id)})


@app.route("/api/unfollow", methods=["POST"])
@login_required
def api_unfollow():
    viewer_id = int(session["user_id"])
    payload = request.get_json(silent=True) or {}
    target_id = int(payload.get("user_id") or 0)
    if not target_id:
        target_id = int(_get_user_id_by_username(str(payload.get("username") or "").strip()) or 0)
    if not target_id:
        return jsonify({"error": "User not found"}), 404
    if target_id == viewer_id:
        return jsonify({"error": "Cannot unfollow yourself"}), 400
    set_follow(viewer_id, target_id, False)
    return jsonify({"ok": True, "is_following": False, "counts": follow_counts(target_id)})


@app.route("/api/conversations", methods=["GET"])
@login_required
def api_conversations():
    user_id = int(session["user_id"])
    return jsonify({"conversations": list_conversations_for_user(user_id)})


@app.route("/api/ai/conversations", methods=["GET"])
@login_required
def api_ai_conversations():
    user_id = int(session["user_id"])
    q = str(request.args.get("q", "") or "").strip()
    try:
        limit = int(request.args.get("limit", "60"))
    except ValueError:
        limit = 60
    return jsonify({"conversations": ai_list_conversations(user_id, q=q, limit=max(1, min(limit, 200)))})


@app.route("/api/ai/conversations", methods=["POST"])
@login_required
def api_ai_create_conversation():
    user_id = int(session["user_id"])
    payload = request.get_json(silent=True) or {}
    title = str(payload.get("title") or "").strip() or "New chat"
    mode = str(payload.get("mode") or "normal").strip().lower()
    system_prompt = str(payload.get("system_prompt") or "").strip()
    conv_id = ai_create_conversation(user_id, title=title, mode=mode, system_prompt=system_prompt)
    return jsonify({"conversation": ai_conversation_brief(conv_id, user_id)})


@app.route("/api/ai/conversations/<int:conversation_id>", methods=["PATCH"])
@login_required
def api_ai_update_conversation(conversation_id: int):
    user_id = int(session["user_id"])
    payload = request.get_json(silent=True) or {}
    title = payload.get("title", None)
    mode = payload.get("mode", None)
    system_prompt = payload.get("system_prompt", None)
    is_pinned = payload.get("is_pinned", None)
    is_bookmarked = payload.get("is_bookmarked", None)
    summary = payload.get("summary", None)
    updated = ai_update_conversation(
        conversation_id,
        user_id,
        title=(None if title is None else str(title)),
        mode=(None if mode is None else str(mode)),
        system_prompt=(None if system_prompt is None else str(system_prompt)),
        is_pinned=(None if is_pinned is None else bool(is_pinned)),
        is_bookmarked=(None if is_bookmarked is None else bool(is_bookmarked)),
    )
    if updated and summary is not None:
        conn = get_conn()
        try:
            cur = conn.cursor()
            cur.execute(
                "UPDATE ai_conversations SET summary = ?, updated_at_utc = ? WHERE id = ? AND user_id = ?",
                (str(summary or ""), utc_now_iso(), conversation_id, user_id),
            )
            conn.commit()
        finally:
            conn.close()
        updated = ai_conversation_brief(conversation_id, user_id)
    if not updated:
        return jsonify({"error": "Not found"}), 404
    return jsonify({"conversation": updated})


@app.route("/api/ai/conversations/<int:conversation_id>", methods=["DELETE"])
@login_required
def api_ai_delete_conversation(conversation_id: int):
    user_id = int(session["user_id"])
    ok = ai_delete_conversation(conversation_id, user_id)
    if not ok:
        return jsonify({"error": "Not found"}), 404
    return jsonify({"ok": True})


@app.route("/api/ai/conversations/<int:conversation_id>/messages", methods=["GET"])
@login_required
def api_ai_messages(conversation_id: int):
    user_id = int(session["user_id"])
    try:
        limit = int(request.args.get("limit", "250"))
    except ValueError:
        limit = 250
    return jsonify({"messages": ai_list_messages(conversation_id, user_id, limit=max(1, min(limit, 500)))})


@app.route("/api/ai/messages/<int:message_id>", methods=["PATCH"])
@login_required
def api_ai_edit_message(message_id: int):
    user_id = int(session["user_id"])
    payload = request.get_json(silent=True) or {}
    content = str(payload.get("content") or "").strip()
    updated = ai_edit_user_message(message_id, user_id, content)
    if not updated:
        return jsonify({"error": "Not found"}), 404
    return jsonify({"ok": True, **updated})


@app.route("/api/ai/memories", methods=["GET"])
@login_required
def api_ai_memories_list():
    user_id = int(session["user_id"])
    try:
        limit = int(request.args.get("limit", "30"))
    except ValueError:
        limit = 30
    return jsonify({"memories": list_user_memories(user_id, limit=max(1, min(limit, 200)))})


@app.route("/api/ai/memories", methods=["POST"])
@login_required
def api_ai_memories_add():
    user_id = int(session["user_id"])
    payload = request.get_json(silent=True) or {}
    text = str(payload.get("text") or "").strip()
    source_conversation_id = payload.get("source_conversation_id", None)
    scid = None
    try:
        if source_conversation_id is not None and str(source_conversation_id).strip():
            scid = int(source_conversation_id)
    except Exception:
        scid = None
    mem = add_user_memory(user_id, text, source_conversation_id=scid)
    if not mem:
        return jsonify({"error": "Empty memory"}), 400
    return jsonify({"memory": mem})


@app.route("/api/ai/memories/<int:memory_id>", methods=["DELETE"])
@login_required
def api_ai_memories_delete(memory_id: int):
    user_id = int(session["user_id"])
    ok = delete_user_memory(user_id, memory_id)
    if not ok:
        return jsonify({"error": "Not found"}), 404
    return jsonify({"ok": True})


@app.route("/api/ai/conversations/<int:conversation_id>/notes", methods=["GET"])
@login_required
def api_ai_notes_get(conversation_id: int):
    user_id = int(session["user_id"])
    note = get_ai_notes(conversation_id, user_id)
    if note is None:
        return jsonify({"error": "Not found"}), 404
    return jsonify({"note": note})


@app.route("/api/ai/conversations/<int:conversation_id>/notes", methods=["PUT"])
@login_required
def api_ai_notes_put(conversation_id: int):
    user_id = int(session["user_id"])
    payload = request.get_json(silent=True) or {}
    body = str(payload.get("body") or "")
    note = upsert_ai_notes(conversation_id, user_id, body)
    if note is None:
        return jsonify({"error": "Not found"}), 404
    return jsonify({"note": note})


@app.route("/api/ai/prompts", methods=["GET"])
@login_required
def api_ai_prompts_list():
    user_id = int(session["user_id"])
    try:
        limit = int(request.args.get("limit", "120"))
    except ValueError:
        limit = 120
    return jsonify({"prompts": list_ai_prompts(user_id, limit=max(1, min(limit, 200)))})


@app.route("/api/ai/prompts", methods=["POST"])
@login_required
def api_ai_prompts_create():
    user_id = int(session["user_id"])
    payload = request.get_json(silent=True) or {}
    title = str(payload.get("title") or "").strip()
    body = str(payload.get("body") or "").strip()
    p = upsert_ai_prompt(user_id, title, body, prompt_id=None)
    if not p:
        return jsonify({"error": "Missing title/body"}), 400
    return jsonify({"prompt": p})


@app.route("/api/ai/prompts/<int:prompt_id>", methods=["PATCH"])
@login_required
def api_ai_prompts_update(prompt_id: int):
    user_id = int(session["user_id"])
    payload = request.get_json(silent=True) or {}
    title = str(payload.get("title") or "").strip()
    body = str(payload.get("body") or "").strip()
    p = upsert_ai_prompt(user_id, title, body, prompt_id=prompt_id)
    if not p:
        return jsonify({"error": "Not found"}), 404
    return jsonify({"prompt": p})


@app.route("/api/ai/prompts/<int:prompt_id>", methods=["DELETE"])
@login_required
def api_ai_prompts_delete(prompt_id: int):
    user_id = int(session["user_id"])
    ok = delete_ai_prompt(user_id, prompt_id)
    if not ok:
        return jsonify({"error": "Not found"}), 404
    return jsonify({"ok": True})


@app.route("/api/ai/conversations/<int:conversation_id>/share", methods=["POST"])
@login_required
def api_ai_share_create(conversation_id: int):
    user_id = int(session["user_id"])
    share = ai_create_or_rotate_share(conversation_id, user_id)
    if share is None:
        return jsonify({"error": "Not found"}), 404
    base = _best_public_base_url()
    if base:
        url = base.rstrip("/") + "/share/" + share["token"]
    else:
        url = url_for("share_page", token=share["token"], _external=True)
    return jsonify({"share": {**share, "url": url}})


@app.route("/api/ai/conversations/<int:conversation_id>/share", methods=["DELETE"])
@login_required
def api_ai_share_revoke(conversation_id: int):
    user_id = int(session["user_id"])
    ok = ai_revoke_share(conversation_id, user_id)
    if not ok:
        return jsonify({"error": "Not found"}), 404
    return jsonify({"ok": True})


def _reserve_ai_quota(user_id: int, requested_tz: str | None) -> tuple[bool, dict[str, Any], bool]:
    """
    Returns (ok, quota_payload, reserved).
    If reserved==True and later work fails, caller should release_daily_message().
    """
    plan_hint = get_user_plan(user_id)
    if plan_hint in PRO_PLANS:
        return True, {"plan": plan_hint}, False

    if requested_tz:
        set_user_tz_if_empty(user_id, str(requested_tz))

    firebase_uid = str(session.get("firebase_uid") or "").strip()
    email = str(session.get("email") or "").strip().lower()
    if firebase_uid and _billing_firestore_client() is not None:
        eff_limit = effective_free_daily_limit_for_user(int(user_id))
        ok2, quota2, reserved2 = _billing_firestore_reserve_daily_message(
            firebase_uid,
            email,
            plan_hint=plan_hint,
            requested_tz=requested_tz,
            free_daily_limit=eff_limit,
            user_id=int(user_id),
        )
        # If Firestore reservation fails open (no limit/used fields), fall back to SQLite
        # so Free plan enforcement still works.
        if (
            plan_hint not in PRO_PLANS
            and int(eff_limit) > 0
            and (not isinstance(quota2, dict) or not isinstance(quota2.get("limit", None), int) or not isinstance(quota2.get("used", None), int))
        ):
            # Continue to SQLite fallback below.
            pass
        else:
            return ok2, quota2, reserved2

    # SQLite fallback.
    plan = FREE_PLAN
    user_tz = get_user_tz(user_id) or normalize_tz_name(str(requested_tz)) or normalize_tz_name((os.getenv("USAGE_TZ") or "").strip())
    day = usage_day_iso(user_tz)
    eff_limit = effective_free_daily_limit_for_user(int(user_id))
    if eff_limit <= 0:
        return True, {"plan": plan, "day": day, "tz": user_tz or ""}, False
    used = get_daily_message_count(user_id, day)
    if used >= eff_limit:
        return (
            False,
            {
                "plan": plan,
                "limit": eff_limit,
                "used": used,
                "day": day,
                "tz": user_tz or "",
                "reset_at": next_reset_iso(user_tz),
            },
            False,
        )
    new_count = reserve_daily_message(user_id, day)
    if new_count > eff_limit:
        release_daily_message(user_id, day)
        return (
            False,
            {
                "plan": plan,
                "limit": eff_limit,
                "used": eff_limit,
                "day": day,
                "tz": user_tz or "",
                "reset_at": next_reset_iso(user_tz),
            },
            False,
        )
    return (
        True,
        {
            "plan": plan,
            "limit": eff_limit,
            "used": new_count,
            "day": day,
            "tz": user_tz or "",
            "reset_at": next_reset_iso(user_tz),
            "store": "sqlite",
        },
        True,
    )


def _release_ai_quota_if_reserved(user_id: int, quota: dict[str, Any]) -> None:
    day = str((quota or {}).get("day") or "").strip()
    if not day:
        return
    # SQLite release is safe even if this quota was reserved in Firestore (row may not exist).
    try:
        release_daily_message(int(user_id), day)
    except Exception:
        pass

    if str((quota or {}).get("store") or "").strip().lower() == "firestore":
        firebase_uid = str(session.get("firebase_uid") or "").strip()
        if firebase_uid:
            _billing_firestore_release_daily_message(firebase_uid, day)


def billing_status_for_user(user_id: int, requested_tz: str | None) -> dict[str, Any]:
    """
    Single source of truth for UI.
    - If Firestore billing doc exists (and Firebase Admin is configured), use it.
    - Otherwise, fall back to SQLite (local dev / legacy).
    """
    user_id = int(user_id)
    plan_hint = get_user_plan(user_id)
    tz = get_user_tz(user_id) or normalize_tz_name(str(requested_tz)) or normalize_tz_name((os.getenv("USAGE_TZ") or "").strip())
    day = usage_day_iso(tz)
    eff_limit = effective_free_daily_limit_for_user(int(user_id))

    firebase_uid = str(session.get("firebase_uid") or "").strip()
    email = str(session.get("email") or "").strip().lower()

    if firebase_uid and _billing_firestore_client() is not None:
        doc = _billing_firestore_ensure_user(firebase_uid, email, user_id=user_id, plan_hint=plan_hint) or {}
        plan = _billing_firestore_effective_plan(doc, plan_hint)
        last_reset = str(doc.get("last_reset_date") or "").strip()
        used = int(doc.get("messages_used_today") or 0)
        used_today = used if last_reset == day else 0
        # If we had to fall back to SQLite for quota enforcement, reflect that usage too.
        if plan not in PRO_PLANS and int(eff_limit) > 0:
            try:
                used_today = max(int(used_today), int(get_daily_message_count(user_id, day)))
            except Exception:
                pass

        return {
            "plan": plan,
            "day": day,
            "tz": tz or "",
            "reset_at": next_reset_iso(tz),
            "free_daily_limit": int(eff_limit),
            "messages_used_today": used_today,
            "subscription_status": str(doc.get("subscription_status") or "").strip(),
            "current_period_end_utc": str(doc.get("current_period_end_utc") or "").strip(),
            "cancel_at_period_end": bool(doc.get("cancel_at_period_end") or False),
            "stripe_customer_id": str(doc.get("stripe_customer_id") or "").strip(),
            "stripe_subscription_id": str(doc.get("stripe_subscription_id") or "").strip(),
        }

    # SQLite fallback.
    plan = plan_hint if plan_hint in {FREE_PLAN, *PRO_PLANS} else FREE_PLAN
    used_today = 0
    if plan not in PRO_PLANS and int(eff_limit) > 0:
        used_today = int(get_daily_message_count(user_id, day))

    sub = _billing_get_subscription(user_id) or {}
    return {
        "plan": plan,
        "day": day,
        "tz": tz or "",
        "reset_at": next_reset_iso(tz),
        "free_daily_limit": int(eff_limit),
        "messages_used_today": used_today,
        "subscription_status": str(sub.get("status") or "").strip(),
        "current_period_end_utc": str(sub.get("current_period_end_utc") or "").strip(),
        "cancel_at_period_end": bool(sub.get("cancel_at_period_end") or False),
        "stripe_customer_id": str(sub.get("customer_id") or "").strip() or _user_get_stripe_customer_id(user_id),
        "stripe_subscription_id": str(sub.get("subscription_id") or "").strip(),
    }


@app.route("/api/ai/conversations/<int:conversation_id>/stream", methods=["POST"])
@login_required
def api_ai_stream(conversation_id: int):
    if not request.is_json:
        return jsonify({"error": "Request must be JSON."}), 400
    user_id = int(session["user_id"])
    if not _ai_conv_belongs_to_user(conversation_id, user_id):
        return jsonify({"error": "Not found"}), 404

    payload = request.get_json(silent=True) or {}
    message = str(payload.get("message") or "").strip()
    mode = str(payload.get("mode") or "").strip().lower()
    regenerate = bool(payload.get("regenerate") or False)
    requested_tz = str(payload.get("tz") or "").strip() or None

    ok, quota, reserved = _reserve_ai_quota(user_id, requested_tz)
    if not ok:
        return jsonify({"error": "Daily free limit reached. Upgrade to Pro.", "quota": quota}), 429

    # Regenerate: delete last assistant message if present and reuse last user message.
    if regenerate:
        conn = get_conn()
        try:
            cur = conn.cursor()
            cur.execute(
                """
                SELECT id, role, content
                FROM ai_messages
                WHERE conversation_id = ?
                ORDER BY id DESC
                LIMIT 2
                """,
                (conversation_id,),
            )
            rows = cur.fetchall()
        finally:
            conn.close()
        if rows and rows[0][1] == "assistant":
            ai_delete_message_and_after(conversation_id, user_id, int(rows[0][0]))
        if rows:
            # If last remaining is user, reuse it.
            if rows[0][1] == "user":
                message = str(rows[0][2] or "").strip()
            elif len(rows) > 1 and rows[1][1] == "user":
                message = str(rows[1][2] or "").strip()
        if not message:
            if reserved:
                _release_ai_quota_if_reserved(user_id, quota)
            return jsonify({"error": "Nothing to regenerate."}), 400

    if not regenerate:
        inserted_user_message_id: int | None = None
        if not message:
            if reserved:
                _release_ai_quota_if_reserved(user_id, quota)
            return jsonify({"error": "Message cannot be empty."}), 400
        if mode and mode in CHAT_MODES:
            ai_update_conversation(conversation_id, user_id, mode=mode)
        inserted_user_message_id = ai_append_message(conversation_id, user_id, "user", message)

        # Auto-title: first real user message.
        brief = ai_conversation_brief(conversation_id, user_id) or {}
        current_title = str(brief.get("title") or "").strip()
        if current_title in {"", "New chat"}:
            title = re.sub(r"\s+", " ", message).strip()[:60]
            if title:
                ai_update_conversation(conversation_id, user_id, title=title)

    else:
        inserted_user_message_id = None

    special = canned_about_app_reply(message)
    if special:
        def gen_special():
            try:
                yield json.dumps({"type": "meta", "mode": (mode or "normal"), "quota": quota}, ensure_ascii=True) + "\n"
                yield json.dumps({"type": "delta", "delta": special}, ensure_ascii=True) + "\n"
                mid = ai_append_message(conversation_id, user_id, "assistant", special)
                msg_obj = {
                    "id": mid,
                    "conversation_id": conversation_id,
                    "role": "assistant",
                    "content": special,
                    "created_at_utc": utc_now_iso(),
                    "edited_at_utc": "",
                }
                yield json.dumps({"type": "done", "assistant": msg_obj, "conversation": ai_conversation_brief(conversation_id, user_id)}, ensure_ascii=True) + "\n"
            except Exception as exc:
                if reserved:
                    _release_ai_quota_if_reserved(user_id, quota)
                if inserted_user_message_id is not None:
                    try:
                        ai_delete_message_and_after(conversation_id, user_id, inserted_user_message_id)
                    except Exception:
                        pass
                yield json.dumps({"type": "error", "error": f"Special reply failed: {str(exc)}"}, ensure_ascii=True) + "\n"

        resp = Response(gen_special(), mimetype="application/x-ndjson")
        resp.headers["Cache-Control"] = "no-cache"
        resp.headers["X-Accel-Buffering"] = "no"
        return resp

    current_client = get_openai_client()
    if current_client is None:
        if reserved:
            _release_ai_quota_if_reserved(user_id, quota)
        return jsonify({"error": "OPENAI_API_KEY is not set on the server."}), 500

    def gen():
        acc = ""
        try:
            mode_used, msgs = ai_build_messages_for_model(conversation_id, user_id)
            stream = current_client.chat.completions.create(
                model=DEFAULT_AI_MODEL,
                messages=msgs,
                stream=True,
            )
            yield json.dumps({"type": "meta", "mode": mode_used, "quota": quota}, ensure_ascii=True) + "\n"
            for ev in stream:
                try:
                    delta = ev.choices[0].delta.content or ""
                except Exception:
                    delta = ""
                if not delta:
                    continue
                acc += delta
                yield json.dumps({"type": "delta", "delta": delta}, ensure_ascii=True) + "\n"
            if not acc.strip():
                acc = "(No response text returned)"
            mid = ai_append_message(conversation_id, user_id, "assistant", acc)
            msg_obj = {
                "id": mid,
                "conversation_id": conversation_id,
                "role": "assistant",
                "content": acc,
                "created_at_utc": utc_now_iso(),
                "edited_at_utc": "",
            }
            yield json.dumps({"type": "done", "assistant": msg_obj, "conversation": ai_conversation_brief(conversation_id, user_id)}, ensure_ascii=True) + "\n"
        except Exception as exc:
            if reserved:
                _release_ai_quota_if_reserved(user_id, quota)
            if inserted_user_message_id is not None:
                # Roll back the user turn so the history doesn't get stuck in a "half sent" state.
                try:
                    ai_delete_message_and_after(conversation_id, user_id, inserted_user_message_id)
                except Exception:
                    pass
            app.logger.error("AI stream failed: %s\n%s", exc, traceback.format_exc())
            yield json.dumps({"type": "error", "error": f"OpenAI request failed: {str(exc)}"}, ensure_ascii=True) + "\n"

    resp = Response(gen(), mimetype="application/x-ndjson")
    resp.headers["Cache-Control"] = "no-cache"
    resp.headers["X-Accel-Buffering"] = "no"
    return resp


def _guess_mime(filename: str) -> str:
    name = (filename or "").lower()
    if name.endswith(".png"):
        return "image/png"
    if name.endswith(".jpg") or name.endswith(".jpeg"):
        return "image/jpeg"
    if name.endswith(".webp"):
        return "image/webp"
    if name.endswith(".gif"):
        return "image/gif"
    if name.endswith(".pdf"):
        return "application/pdf"
    if name.endswith(".md"):
        return "text/markdown"
    if name.endswith(".txt"):
        return "text/plain"
    if name.endswith(".csv"):
        return "text/csv"
    if name.endswith(".json"):
        return "application/json"
    if name.endswith(".docx"):
        return "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    return "application/octet-stream"


def _read_textish_file(path: str, mime: str) -> str:
    # Keep this safe and dependency-light. PDF/DOCX are optional.
    if mime in {"text/plain", "text/markdown", "text/csv", "application/json"}:
        raw = open(path, "rb").read()
        # Best-effort: utf-8 then latin-1.
        try:
            return raw.decode("utf-8")
        except Exception:
            return raw.decode("latin-1", errors="replace")
    if mime == "application/pdf":
        try:
            from pypdf import PdfReader  # type: ignore
        except Exception as exc:
            raise RuntimeError("PDF support requires 'pypdf'. Install it to enable PDF summarization.") from exc
        reader = PdfReader(path)
        chunks = []
        for p in reader.pages[:30]:
            try:
                chunks.append(p.extract_text() or "")
            except Exception:
                pass
        return "\n\n".join([c for c in chunks if c.strip()]).strip()
    if mime.endswith("wordprocessingml.document"):
        try:
            import docx  # type: ignore
        except Exception as exc:
            raise RuntimeError("DOCX support requires 'python-docx'. Install it to enable DOCX summarization.") from exc
        d = docx.Document(path)
        return "\n".join([p.text for p in d.paragraphs]).strip()
    raise RuntimeError("Unsupported file type for text extraction.")


@app.route("/api/ai/conversations/<int:conversation_id>/files", methods=["GET"])
@login_required
def api_ai_files_list(conversation_id: int):
    user_id = int(session["user_id"])
    if not _ai_conv_belongs_to_user(conversation_id, user_id):
        return jsonify({"error": "Not found"}), 404
    conn = get_conn()
    try:
        cur = conn.cursor()
        cur.execute(
            """
            SELECT id, filename, mime, size_bytes, storage_path, created_at_utc
            FROM ai_files
            WHERE conversation_id = ? AND user_id = ?
            ORDER BY id DESC
            LIMIT 50
            """,
            (conversation_id, user_id),
        )
        files = [
            {
                "id": int(r[0]),
                "filename": r[1],
                "mime": r[2],
                "size_bytes": int(r[3] or 0),
                "storage_path": r[4],
                "url": url_for("static", filename=r[4].replace("\\", "/")),
                "created_at_utc": r[5],
            }
            for r in cur.fetchall()
        ]
    finally:
        conn.close()
    return jsonify({"files": files})


@app.route("/api/ai/conversations/<int:conversation_id>/files", methods=["POST"])
@login_required
def api_ai_files_upload(conversation_id: int):
    user_id = int(session["user_id"])
    if not _ai_conv_belongs_to_user(conversation_id, user_id):
        return jsonify({"error": "Not found"}), 404
    if "file" not in request.files:
        return jsonify({"error": "Missing file"}), 400
    f = request.files["file"]
    filename = secure_filename(f.filename or "")
    if not filename:
        return jsonify({"error": "Invalid filename"}), 400

    # Enforce size limit without reading full file into memory if possible.
    try:
        f.stream.seek(0, os.SEEK_END)
        size = int(f.stream.tell() or 0)
        f.stream.seek(0)
    except Exception:
        size = 0
    if size and size > AI_FILE_MAX_BYTES:
        return jsonify({"error": f"File too large (max {AI_FILE_MAX_BYTES} bytes)"}), 413

    mime = _guess_mime(filename)
    ext = os.path.splitext(filename)[1].lower().lstrip(".")
    allowed = set(ALLOWED_IMAGE_EXTENSIONS) | {"pdf", "txt", "md", "csv", "json", "docx"}
    if ext and ext not in allowed:
        return jsonify({"error": f"Unsupported file type: .{ext}"}), 400

    # Store under static/uploads/ai/u_<id>/c_<id>/
    rel_dir = os.path.join("uploads", "ai", f"u_{user_id}", f"c_{conversation_id}")
    abs_dir = _ensure_user_dir("ai", f"u_{user_id}", f"c_{conversation_id}")
    stored_name = f"{uuid4().hex}_{filename}"
    abs_path = os.path.join(abs_dir, stored_name)
    f.save(abs_path)

    # Determine size after save.
    try:
        size = int(os.path.getsize(abs_path))
    except Exception:
        size = 0

    now = _safe_now_iso()
    storage_path = os.path.join(rel_dir, stored_name)
    conn = get_conn()
    try:
        cur = conn.cursor()
        cur.execute(
            """
            INSERT INTO ai_files (conversation_id, user_id, filename, mime, size_bytes, storage_path, created_at_utc)
            VALUES (?, ?, ?, ?, ?, ?, ?)
            """,
            (conversation_id, user_id, filename, mime, size, storage_path, now),
        )
        file_id = int(cur.lastrowid)
        conn.commit()
    finally:
        conn.close()
    return jsonify(
        {
            "file": {
                "id": file_id,
                "filename": filename,
                "mime": mime,
                "size_bytes": size,
                "storage_path": storage_path,
                "url": url_for("static", filename=storage_path.replace("\\", "/")),
                "created_at_utc": now,
            }
        }
    )


@app.route("/api/ai/files/<int:file_id>", methods=["DELETE"])
@login_required
def api_ai_files_delete(file_id: int):
    user_id = int(session["user_id"])
    conn = get_conn()
    try:
        cur = conn.cursor()
        cur.execute(
            "SELECT storage_path FROM ai_files WHERE id = ? AND user_id = ?",
            (int(file_id), user_id),
        )
        row = cur.fetchone()
        if not row:
            return jsonify({"error": "Not found"}), 404
        storage_path = row[0] or ""
        cur.execute("DELETE FROM ai_files WHERE id = ? AND user_id = ?", (int(file_id), user_id))
        conn.commit()
    finally:
        conn.close()
    # Best-effort delete the stored file.
    try:
        abs_path = os.path.join(BASE_DIR, "static", storage_path)
        if abs_path.startswith(os.path.join(BASE_DIR, "static")) and os.path.exists(abs_path):
            os.remove(abs_path)
    except Exception:
        pass
    return jsonify({"ok": True})


@app.route("/api/ai/tools/improve-prompt", methods=["POST"])
@login_required
def api_ai_tool_improve_prompt():
    if not request.is_json:
        return jsonify({"error": "Request must be JSON."}), 400
    user_id = int(session["user_id"])
    payload = request.get_json(silent=True) or {}
    draft = str(payload.get("draft") or "").strip()
    if not draft:
        return jsonify({"error": "Empty draft"}), 400
    current_client = get_openai_client()
    if current_client is None:
        return jsonify({"error": "OPENAI_API_KEY is not set on the server."}), 500
    # Quota: treat as one message.
    ok, quota, reserved = _reserve_ai_quota(user_id, str(payload.get("tz") or "").strip() or None)
    if not ok:
        return jsonify({"error": "Daily free limit reached. Upgrade to Pro.", "quota": quota}), 429
    try:
        msgs = [
            {
                "role": "system",
                "content": "Rewrite the user's draft into a clearer, more specific prompt for an AI assistant. Keep intent. Add constraints, context, and desired output format. Return only the improved prompt.",
            },
            {"role": "user", "content": draft},
        ]
        resp = current_client.chat.completions.create(model=DEFAULT_AI_MODEL, messages=msgs)
        out = (resp.choices[0].message.content or "").strip()
        if not out:
            out = draft
        return jsonify({"prompt": out, "quota": quota})
    except Exception as exc:
        if reserved:
            _release_ai_quota_if_reserved(user_id, quota)
        return jsonify({"error": f"OpenAI request failed: {str(exc)}"}), 500


@app.route("/api/ai/tools/summarize-conversation/<int:conversation_id>", methods=["POST"])
@login_required
def api_ai_tool_summarize_conversation(conversation_id: int):
    if not request.is_json:
        return jsonify({"error": "Request must be JSON."}), 400
    user_id = int(session["user_id"])
    if not _ai_conv_belongs_to_user(conversation_id, user_id):
        return jsonify({"error": "Not found"}), 404
    current_client = get_openai_client()
    if current_client is None:
        return jsonify({"error": "OPENAI_API_KEY is not set on the server."}), 500
    payload = request.get_json(silent=True) or {}
    ok, quota, reserved = _reserve_ai_quota(user_id, str(payload.get("tz") or "").strip() or None)
    if not ok:
        return jsonify({"error": "Daily free limit reached. Upgrade to Pro.", "quota": quota}), 429
    try:
        _, msgs = ai_build_messages_for_model(conversation_id, user_id)
        # Summarize only the conversation turns (exclude system).
        convo = [m for m in msgs if m.get("role") in {"user", "assistant"}]
        prompt = (
            "Summarize this conversation into:\n"
            "1) Key goals\n"
            "2) Decisions / conclusions\n"
            "3) Open questions\n"
            "4) Next actions\n"
            "Keep it concise and factual."
        )
        resp = current_client.chat.completions.create(
            model=DEFAULT_AI_MODEL,
            messages=[{"role": "system", "content": prompt}, {"role": "user", "content": json.dumps(convo, ensure_ascii=True)}],
        )
        summary = (resp.choices[0].message.content or "").strip()
        if not summary:
            summary = "(empty summary)"
        ai_update_conversation(conversation_id, user_id, title=None)
        conn = get_conn()
        try:
            cur = conn.cursor()
            cur.execute(
                "UPDATE ai_conversations SET summary = ?, updated_at_utc = ? WHERE id = ? AND user_id = ?",
                (summary, utc_now_iso(), conversation_id, user_id),
            )
            conn.commit()
        finally:
            conn.close()
        return jsonify({"summary": summary, "quota": quota, "conversation": ai_conversation_brief(conversation_id, user_id)})
    except Exception as exc:
        if reserved:
            _release_ai_quota_if_reserved(user_id, quota)
        return jsonify({"error": f"OpenAI request failed: {str(exc)}"}), 500


@app.route("/api/ai/tools/analyze-file", methods=["POST"])
@login_required
def api_ai_tool_analyze_file():
    """
    Analyze an uploaded file with AI. For text/PDF/DOCX we extract text; for images we require OPENAI_VISION_MODEL.
    """
    if not request.is_json:
        return jsonify({"error": "Request must be JSON."}), 400
    user_id = int(session["user_id"])
    payload = request.get_json(silent=True) or {}
    file_id = int(payload.get("file_id") or 0)
    question = str(payload.get("question") or "Summarize this file.").strip()
    if not file_id:
        return jsonify({"error": "Missing file_id"}), 400
    current_client = get_openai_client()
    if current_client is None:
        return jsonify({"error": "OPENAI_API_KEY is not set on the server."}), 500
    ok, quota, reserved = _reserve_ai_quota(user_id, str(payload.get("tz") or "").strip() or None)
    if not ok:
        return jsonify({"error": "Daily free limit reached. Upgrade to Pro.", "quota": quota}), 429

    conn = get_conn()
    try:
        cur = conn.cursor()
        cur.execute(
            "SELECT conversation_id, filename, mime, storage_path FROM ai_files WHERE id = ? AND user_id = ?",
            (file_id, user_id),
        )
        row = cur.fetchone()
    finally:
        conn.close()
    if not row:
        if reserved:
            _release_ai_quota_if_reserved(user_id, quota)
        return jsonify({"error": "File not found"}), 404

    conv_id = int(row[0])
    filename = row[1] or ""
    mime = row[2] or "application/octet-stream"
    storage_path = row[3] or ""
    abs_path = os.path.join(BASE_DIR, "static", storage_path)

    try:
        if mime.startswith("image/"):
            if not DEFAULT_AI_VISION_MODEL:
                raise RuntimeError("Image analysis requires OPENAI_VISION_MODEL (e.g. gpt-4o-mini).")
            import base64

            raw = open(abs_path, "rb").read()
            b64 = base64.b64encode(raw).decode("ascii")
            data_url = f"data:{mime};base64,{b64}"
            msgs = [
                {
                    "role": "system",
                    "content": "You are an expert visual analyst. Answer the user's question about the image accurately and concisely.",
                },
                {
                    "role": "user",
                    "content": [
                        {"type": "text", "text": f"Question: {question}\n\nFilename: {filename}"},
                        {"type": "image_url", "image_url": {"url": data_url}},
                    ],
                },
            ]
            resp = current_client.chat.completions.create(model=DEFAULT_AI_VISION_MODEL, messages=msgs)
            out = (resp.choices[0].message.content or "").strip()
        else:
            extracted = _read_textish_file(abs_path, mime)
            if not extracted:
                extracted = "(No extractable text found.)"
            msgs = [
                {"role": "system", "content": "You are a document analyst. Use the provided document text to answer the question."},
                {"role": "user", "content": f"Question: {question}\n\nDocument text:\n{extracted[:120000]}"},
            ]
            resp = current_client.chat.completions.create(model=DEFAULT_AI_MODEL, messages=msgs)
            out = (resp.choices[0].message.content or "").strip()

        if not out:
            out = "(empty response)"
        # Optionally drop the result into the conversation.
        if conv_id and _ai_conv_belongs_to_user(conv_id, user_id) and bool(payload.get("append_to_chat") or False):
            ai_append_message(conv_id, user_id, "assistant", out)
        return jsonify({"result": out, "quota": quota})
    except Exception as exc:
        if reserved:
            _release_ai_quota_if_reserved(user_id, quota)
        return jsonify({"error": str(exc)}), 500


@app.route("/api/conversations/dm", methods=["POST"])
@login_required
def api_dm_conversation():
    user_id = int(session["user_id"])
    payload = request.get_json(silent=True) or {}
    peer_id = int(payload.get("user_id") or 0)
    if not peer_id:
        return jsonify({"error": "Missing user_id"}), 400
    if peer_id == user_id:
        return jsonify({"error": "Cannot DM yourself"}), 400
    conv_id = ensure_dm_conversation(user_id, peer_id)
    brief = conversation_brief_for_user(conv_id, user_id)
    return jsonify({"conversation": brief})


@app.route("/api/conversations/<int:conversation_id>/messages", methods=["GET"])
@login_required
def api_conversation_messages(conversation_id: int):
    user_id = int(session["user_id"])
    try:
        limit = int(request.args.get("limit", "80"))
    except ValueError:
        limit = 80
    return jsonify({"messages": list_messages(conversation_id, user_id, limit=limit)})


@app.route("/api/conversations/<int:conversation_id>/messages", methods=["POST"])
@login_required
def api_send_message(conversation_id: int):
    user_id = int(session["user_id"])
    if not user_is_member(conversation_id, user_id):
        return jsonify({"error": "Forbidden"}), 403
    payload = request.get_json(silent=True) or {}
    body = str(payload.get("body") or "").strip()
    if not body:
        return jsonify({"error": "Empty message"}), 400
    created_at = utc_now_iso()
    conn = get_conn()
    try:
        cur = conn.cursor()
        cur.execute(
            """
            INSERT INTO chat_messages (conversation_id, sender_id, type, body, media_path, created_at_utc)
            VALUES (?, ?, 'text', ?, '', ?)
            """,
            (conversation_id, user_id, body, created_at),
        )
        msg_id = int(cur.lastrowid)
        conn.commit()
    finally:
        conn.close()

    msg = {
        "id": msg_id,
        "conversation_id": conversation_id,
        "sender_id": user_id,
        "type": "text",
        "body": body,
        "media_path": "",
        "created_at_utc": created_at,
    }
    return jsonify({"message": msg})


@app.route("/api/conversations/<int:conversation_id>/read", methods=["POST"])
@login_required
def api_conversation_read(conversation_id: int):
    user_id = int(session["user_id"])
    mark_conversation_read(conversation_id, user_id)
    brief = conversation_brief_for_user(conversation_id, user_id)
    return jsonify({"ok": True, "conversation": brief})


@app.route("/account")
@login_required
def account_page():
    if session.get("is_guest"):
        return redirect(url_for("auth_page", error="Guest accounts can't open Account. Create an account to continue.", tab="signup"))
    user_id = int(session["user_id"])
    profile = get_user_profile(user_id)
    if not profile:
        return redirect(url_for("auth_page", error="User profile not found."))
    counts = follow_counts(int(user_id))
    billing = billing_status_for_user(int(user_id), requested_tz=None)
    return render_template(
        "account.html",
        profile=profile,
        counts=counts,
        billing=billing,
        message=request.args.get("message", ""),
        error=request.args.get("error", ""),
        firebase_config=_firebase_web_config(),
    )


@app.route("/account/update", methods=["POST"])
@login_required
def account_update():
    require_csrf()
    if session.get("is_guest"):
        return redirect(url_for("auth_page", error="Guest accounts can't update profile. Create an account to continue.", tab="signup"))
    user_id = session["user_id"]
    profile = get_user_profile(user_id)
    if not profile:
        return redirect(url_for("auth_page", error="User profile not found."))
    billing = billing_status_for_user(int(user_id), requested_tz=None)

    username = request.form.get("username", "").strip()
    # Allow partial updates (some tabs/forms may omit fields).
    email_raw = request.form.get("email", None)
    email_in_form = email_raw is not None
    email = (email_raw.strip() if isinstance(email_raw, str) else str(profile.get("email", "") or "").strip())
    address = request.form.get("address", "").strip()
    age_raw = request.form.get("age", "").strip()
    age = None

    if not re.fullmatch(r"[A-Za-z0-9_]{3,24}", username or ""):
        return render_template(
            "account.html",
            profile={**profile, "username": username, "email": email, "address": address, "age": age_raw},
            error="Username must be 3-24 characters: letters, numbers, underscore.",
            message="",
            counts=follow_counts(int(user_id)),
            billing=billing,
            firebase_config=_firebase_web_config(),
        )

    def _looks_like_email(s: str) -> bool:
        s = (s or "").strip()
        return bool(s) and ("@" in s) and (not re.search(r"\\s", s))

    if email_in_form and email and (not _looks_like_email(email)):
        return render_template(
            "account.html",
            profile={**profile, "username": username, "email": email, "address": address, "age": age_raw},
            error="Please enter a valid email address.",
            message="",
            counts=follow_counts(int(user_id)),
            billing=billing,
            firebase_config=_firebase_web_config(),
        )

    if age_raw:
        try:
            age = int(age_raw)
            if age < 0 or age > 120:
                raise ValueError
        except ValueError:
            return render_template(
                "account.html",
                profile={**profile, "username": username, "email": email, "address": address, "age": age_raw},
                error="Age must be a number between 0 and 120.",
                message="",
                counts=follow_counts(int(user_id)),
                billing=billing,
                firebase_config=_firebase_web_config(),
            )

    conn = get_conn()
    cur = conn.cursor()
    cur.execute("SELECT id FROM users WHERE username = ? COLLATE NOCASE AND id != ?", (username, user_id))
    if cur.fetchone():
        conn.close()
        return render_template(
            "account.html",
            profile={**profile, "username": username, "email": email, "address": address, "age": age_raw},
            error="Username is already taken.",
            message="",
            counts=follow_counts(int(user_id)),
            billing=billing,
            firebase_config=_firebase_web_config(),
        )

    if email_in_form and email:
        cur.execute("SELECT id FROM users WHERE email = ? COLLATE NOCASE AND id != ?", (email, user_id))
        if cur.fetchone():
            conn.close()
            return render_template(
                "account.html",
                profile={**profile, "username": username, "email": email, "address": address, "age": age_raw},
                error="Email is already in use.",
                message="",
                counts=follow_counts(int(user_id)),
                billing=billing,
                firebase_config=_firebase_web_config(),
            )

    photo_path = profile.get("photo_path", "")
    try:
        uploaded_photo = save_profile_photo(request.files.get("photo"))
        if uploaded_photo:
            photo_path = uploaded_photo
    except ValueError as exc:
        conn.close()
        return render_template(
            "account.html",
            profile={**profile, "username": username, "email": email, "address": address, "age": age_raw},
            error=str(exc),
            message="",
            counts=follow_counts(int(user_id)),
            billing=billing,
            firebase_config=_firebase_web_config(),
        )

    cur.execute(
        """
        UPDATE users
        SET username = ?, email = ?, address = ?, age = ?, photo_path = ?
        WHERE id = ?
        """,
        (username, email, address, age, photo_path, user_id),
    )
    conn.commit()
    conn.close()

    session["username"] = username
    session["display_name"] = username
    return redirect(url_for("account_page", message="Profile updated successfully."))


@app.route("/account/change-password", methods=["POST"])
@login_required
def account_change_password():
    require_csrf()
    if session.get("is_guest"):
        return redirect(url_for("auth_page", error="Guest accounts don't have a password. Create an account to continue.", tab="signup"))
    return redirect(url_for("account_page", error="Password is managed by Firebase."))


@app.route("/account/add-account", methods=["POST"])
@login_required
def account_add_account():
    require_csrf()
    user_id = session.get("user_id")
    if user_id in conversations:
        conversations.pop(user_id, None)
    session.clear()
    return redirect(url_for("auth_page", error="Logged out. You can create or login with another account."))


@app.route("/feedback", methods=["GET", "POST"])
def feedback_page():
    if request.method == "GET":
        return render_template(
            "feedback.html",
            username=session.get("username", ""),
            message=request.args.get("message", ""),
            error=request.args.get("error", ""),
        )

    require_csrf()
    category = (request.form.get("category", "") or "").strip()
    rating_raw = (request.form.get("rating", "") or "").strip()
    text = (request.form.get("message", "") or "").strip()
    contact_email = (request.form.get("contact_email", "") or "").strip()
    page = (request.form.get("page", "") or "").strip()
    tz = (request.form.get("tz", "") or "").strip()

    rating = None
    if rating_raw:
        try:
            rating = int(rating_raw)
        except ValueError:
            rating = None
    if rating is not None and (rating < 0 or rating > 10):
        rating = None

    if not category:
        return redirect(url_for("feedback_page", error="Please choose a category."))
    if not text or len(text) < 10:
        return redirect(url_for("feedback_page", error="Please write at least 10 characters."))

    user_id = session.get("user_id")
    username = session.get("username")
    user_agent = request.headers.get("User-Agent", "")
    ip = request.headers.get("X-Forwarded-For", "").split(",")[0].strip() or (request.remote_addr or "")

    created_at_utc = datetime.now(timezone.utc).isoformat()
    conn = get_conn()
    try:
        cur = conn.cursor()
        cur.execute(
            """
            INSERT INTO feedback (
                created_at_utc, user_id, username, contact_email, rating, category, message, page, tz, user_agent, ip
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (created_at_utc, user_id, username, contact_email, rating, category, text, page, tz, user_agent, ip),
        )
        conn.commit()
    finally:
        conn.close()

    return redirect(url_for("feedback_page", message="Thanks. Brutal honesty helps."))


@app.route("/robots.txt")
def robots():
    content = "User-agent: *\nAllow: /\nSitemap: " + url_for("sitemap", _external=True) + "\n"
    return Response(content, mimetype="text/plain")


@app.route("/sitemap.xml")
def sitemap():
    pages = [
        url_for("landing", _external=True),
        url_for("auth_page", _external=True),
        url_for("feedback_page", _external=True),
    ]
    xml_items = "".join([f"<url><loc>{p}</loc></url>" for p in pages])
    xml = (
        '<?xml version="1.0" encoding="UTF-8"?>'
        '<urlset xmlns="http://www.sitemaps.org/schemas/sitemap/0.9">'
        f"{xml_items}"
        "</urlset>"
    )
    return Response(xml, mimetype="application/xml")


@app.route("/share/<token>")
def share_page(token: str):
    tok = (token or "").strip()
    if not tok or len(tok) < 16:
        return Response("Not found", status=404)
    conn = get_conn()
    try:
        cur = conn.cursor()
        cur.execute(
            """
            SELECT s.conversation_id, s.revoked_at_utc
            FROM ai_shares s
            WHERE s.token = ?
            ORDER BY s.id DESC
            LIMIT 1
            """,
            (tok,),
        )
        row = cur.fetchone()
        if not row or (row[1] or "").strip():
            return Response("Not found", status=404)
        conv_id = int(row[0])
        cur.execute("SELECT title FROM ai_conversations WHERE id = ?", (conv_id,))
        c = cur.fetchone()
        title = (c[0] if c else "Shared chat") or "Shared chat"
        cur.execute(
            "SELECT role, content, created_at_utc FROM ai_messages WHERE conversation_id = ? ORDER BY id ASC LIMIT 800",
            (conv_id,),
        )
        messages = [{"role": r[0], "content": r[1], "created_at_utc": r[2]} for r in cur.fetchall()]
    finally:
        conn.close()
    return render_template("share.html", title=title, messages=messages)


@app.route("/chat", methods=["POST"])
@login_required
def chat():
    if not request.is_json:
        return jsonify({"error": "Request must be JSON."}), 400

    user_message = request.json.get("message", "").strip()
    if not user_message:
        return jsonify({"error": "Message cannot be empty."}), 400

    user_id = int(session["user_id"])
    requested_tz = None
    try:
        requested_tz = str(request.json.get("tz") or "").strip() or None
    except Exception:
        requested_tz = None

    ok, quota, reserved = _reserve_ai_quota(user_id, requested_tz)
    if not ok:
        return jsonify({"error": "Daily free limit reached. Upgrade to Pro.", "quota": quota}), 429

    special = canned_about_app_reply(user_message)
    if special:
        chat_id = str(request.json.get("chat_id", "default")).strip() or "default"
        mode = str(request.json.get("mode", "normal")).strip().lower()
        if mode not in CHAT_MODES:
            mode = "normal"
        user_chats = get_user_chats(user_id)
        history = user_chats.setdefault(chat_id, [])
        history.append({"role": "user", "content": user_message})
        history.append({"role": "assistant", "content": special})
        return jsonify({"reply": special, "mode": mode, "special": "about_app", "quota": quota})

    current_client = get_openai_client()
    if current_client is None:
        if reserved:
            _release_ai_quota_if_reserved(user_id, quota)
        return jsonify({"error": "OPENAI_API_KEY is not set on the server."}), 500

    plan = str((quota or {}).get("plan") or FREE_PLAN).strip().lower() or FREE_PLAN
    chat_id = str(request.json.get("chat_id", "default")).strip() or "default"
    mode = str(request.json.get("mode", "normal")).strip().lower()
    if mode not in CHAT_MODES:
        mode = "normal"
    user_chats = get_user_chats(user_id)
    history = user_chats.setdefault(chat_id, [])
    history.append({"role": "user", "content": user_message})

    try:
        messages_for_api = [{"role": "system", "content": CHAT_MODES[mode]}] + history
        response = current_client.chat.completions.create(
            model="gpt-4.1-mini",
            messages=messages_for_api,
        )
        reply = response.choices[0].message.content or "(No response text returned)"
        history.append({"role": "assistant", "content": reply})
        payload: dict[str, Any] = {"reply": reply, "mode": mode}
        payload["quota"] = quota or {"plan": plan}
        return jsonify(payload)
    except Exception as exc:
        if reserved:
            _release_ai_quota_if_reserved(user_id, quota)
        if history and history[-1].get("role") == "user" and history[-1].get("content") == user_message:
            history.pop()
        app.logger.error("OpenAI /chat failed: %s\n%s", exc, traceback.format_exc())
        return jsonify({"error": f"OpenAI request failed: {str(exc)}"}), 500


@app.route("/reset", methods=["POST"])
@login_required
def reset_chat():
    user_id = session["user_id"]
    user_chats = get_user_chats(user_id)
    chat_id = "default"
    if request.is_json:
        chat_id = str(request.json.get("chat_id", "default")).strip() or "default"
    user_chats[chat_id] = []
    return jsonify({"ok": True})


@app.route("/admin/set-user-plan", methods=["POST"])
def admin_set_user_plan():
    ok, resp = admin_required()
    if not ok:
        # Keep existing behavior: hidden if admin disabled, forbidden otherwise.
        if (os.getenv("ADMIN_TOKEN") or "").strip():
            return resp, 403
        return resp, 404
    if not request.is_json:
        return jsonify({"error": "Request must be JSON."}), 400

    username = str(request.json.get("username", "")).strip()
    user_id_raw = request.json.get("user_id", None)
    plan = str(request.json.get("plan", "")).strip().lower()
    if not plan:
        return jsonify({"error": "Missing 'plan'."}), 400
    if plan not in {FREE_PLAN, *PRO_PLANS}:
        return jsonify({"error": f"Unsupported plan: {plan}"}), 400

    conn = get_conn()
    try:
        cur = conn.cursor()
        if user_id_raw is not None and str(user_id_raw).strip():
            try:
                user_id = int(user_id_raw)
            except ValueError:
                return jsonify({"error": "user_id must be an integer."}), 400
            cur.execute("UPDATE users SET plan = ? WHERE id = ?", (plan, user_id))
        elif username:
            cur.execute("UPDATE users SET plan = ? WHERE username = ?", (plan, username))
        else:
            return jsonify({"error": "Provide 'user_id' or 'username'."}), 400
        conn.commit()
        if cur.rowcount <= 0:
            return jsonify({"error": "User not found."}), 404
        return jsonify({"ok": True, "plan": plan})
    finally:
        conn.close()


@app.route("/admin/feedback", methods=["GET"])
def admin_feedback():
    ok, resp = admin_required()
    if not ok:
        if (os.getenv("ADMIN_TOKEN") or "").strip():
            return resp, 403
        return resp, 404

    message = (request.args.get("message", "") or "").strip()
    limit_raw = (request.args.get("limit", "") or "").strip()
    try:
        limit = int(limit_raw) if limit_raw else 200
    except ValueError:
        limit = 200
    limit = max(1, min(2000, limit))

    conn = get_conn()
    rows: list[dict[str, Any]] = []
    try:
        cur = conn.cursor()
        cur.execute(
            """
            SELECT id, created_at_utc, COALESCE(username, ''), COALESCE(contact_email, ''), COALESCE(category, ''),
                   COALESCE(rating, ''), COALESCE(message, ''), COALESCE(page, ''), COALESCE(tz, ''),
                   COALESCE(user_agent, ''), COALESCE(ip, '')
            FROM feedback
            ORDER BY id DESC
            LIMIT ?
            """,
            (limit,),
        )
        for r in cur.fetchall():
            rows.append(
                {
                    "id": r[0],
                    "created_at_utc": r[1],
                    "username": r[2],
                    "contact_email": r[3],
                    "category": r[4],
                    "rating": r[5],
                    "message": r[6],
                    "page": r[7],
                    "tz": r[8],
                    "user_agent": r[9],
                    "ip": r[10],
                }
            )
    except sqlite3.OperationalError as exc:
        # Table might not exist yet if init_db hasn't run in this process.
        return render_template(
            "admin_feedback.html",
            rows=[],
            error=f"DB error: {exc}",
            message=message,
            limit=limit,
            token=request.args.get("token", ""),
        )
    finally:
        conn.close()

    return render_template(
        "admin_feedback.html",
        rows=rows,
        error="",
        message=message,
        limit=limit,
        token=request.args.get("token", ""),
    )


@app.route("/admin/users", methods=["GET"])
def admin_users():
    ok, resp = admin_required()
    if not ok:
        if (os.getenv("ADMIN_TOKEN") or "").strip():
            return resp, 403
        return resp, 404

    token = (request.args.get("token", "") or "").strip()
    q = (request.args.get("q", "") or "").strip()
    message = (request.args.get("message", "") or "").strip()
    limit_raw = (request.args.get("limit", "") or "").strip()
    try:
        limit = int(limit_raw) if limit_raw else 500
    except ValueError:
        limit = 500
    limit = max(1, min(5000, limit))

    conn = get_conn()
    users: list[dict[str, Any]] = []
    try:
        cur = conn.cursor()
        needle = q.strip().lower()
        if needle:
            like = f"%{needle}%"
            cur.execute(
                """
                SELECT id, username, COALESCE(email, ''), COALESCE(plan, ''), COALESCE(tz, ''),
                       COALESCE(created_at_utc, ''), COALESCE(last_login_at_utc, ''), COALESCE(last_login_ip, ''),
                       daily_limit_override
                FROM users
                WHERE LOWER(username) LIKE ? OR LOWER(COALESCE(email, '')) LIKE ?
                ORDER BY id DESC
                LIMIT ?
                """,
                (like, like, limit),
            )
        else:
            cur.execute(
                """
                SELECT id, username, COALESCE(email, ''), COALESCE(plan, ''), COALESCE(tz, ''),
                       COALESCE(created_at_utc, ''), COALESCE(last_login_at_utc, ''), COALESCE(last_login_ip, ''),
                       daily_limit_override
                FROM users
                ORDER BY id DESC
                LIMIT ?
                """,
                (limit,),
            )
        rows = cur.fetchall()
        user_ids = [int(r[0]) for r in rows]

        firebase_by_user: dict[int, str] = {}
        if user_ids:
            qs_marks = ",".join(["?"] * len(user_ids))
            cur.execute(
                f"""
                SELECT user_id, provider_user_id
                FROM identities
                WHERE provider = 'firebase' AND user_id IN ({qs_marks})
                """,
                tuple(user_ids),
            )
            for r in cur.fetchall():
                try:
                    firebase_by_user[int(r[0])] = str(r[1] or "").strip()
                except Exception:
                    pass

        usage_tz_default = normalize_tz_name((os.getenv("USAGE_TZ") or "").strip()) or "UTC"
        for r in rows:
            uid = int(r[0])
            username = str(r[1] or "")
            email = str(r[2] or "")
            plan = (str(r[3] or "") or FREE_PLAN).strip().lower() or FREE_PLAN
            tz = normalize_tz_name(str(r[4] or "")) or ""
            created_at_utc = str(r[5] or "")
            last_login_at_utc = str(r[6] or "")
            last_login_ip = str(r[7] or "")
            daily_limit_override = r[8]
            try:
                daily_limit_override = int(daily_limit_override) if daily_limit_override is not None else None
            except Exception:
                daily_limit_override = None

            eff_limit = effective_free_daily_limit_for_user(uid)
            day = usage_day_iso(tz or usage_tz_default)
            used_today = 0
            reset_at = ""
            try:
                reset_at = next_reset_iso(tz or usage_tz_default)
            except Exception:
                reset_at = ""
            try:
                used_today = int(get_daily_message_count(uid, day)) if plan not in PRO_PLANS else 0
            except Exception:
                used_today = 0

            users.append(
                {
                    "id": uid,
                    "username": username,
                    "email": email,
                    "plan": plan,
                    "tz": tz,
                    "created_at_utc": created_at_utc,
                    "last_login_at_utc": last_login_at_utc,
                    "last_login_ip": last_login_ip,
                    "daily_limit_override": daily_limit_override,
                    "effective_limit": eff_limit,
                    "used_today": used_today,
                    "reset_at": reset_at,
                    "firebase_uid": firebase_by_user.get(uid, ""),
                }
            )
    finally:
        conn.close()

    return render_template(
        "admin_users.html",
        users=users,
        q=q,
        message=message,
        token=token,
        free_daily_limit=int(FREE_DAILY_MESSAGE_LIMIT),
        pro_plans=PRO_PLANS,
        now_utc=utc_now_iso(),
    )


@app.route("/admin/users/<int:user_id>/plan", methods=["POST"])
def admin_user_set_plan(user_id: int):
    ok, resp = admin_required()
    if not ok:
        if (os.getenv("ADMIN_TOKEN") or "").strip():
            return resp, 403
        return resp, 404

    token = (request.args.get("token", "") or "").strip()
    plan = str(request.form.get("plan") or "").strip().lower()
    if plan not in {FREE_PLAN, "pro", "ultimate"}:
        return redirect(url_for("admin_users", token=token, message="Unsupported plan."))
    desired = FREE_PLAN if plan == FREE_PLAN else ("ultimate" if plan == "ultimate" else "pro")

    try:
        _set_user_plan(int(user_id), desired)
    except Exception:
        return redirect(url_for("admin_users", token=token, message="Failed to update plan."))

    try:
        firebase_uid = _firebase_uid_for_user_id(int(user_id))
        if firebase_uid:
            doc_ref = _billing_firestore_doc(firebase_uid)
            if doc_ref is not None:
                doc_ref.set({"uid": firebase_uid, "plan_type": desired, "updated_at_utc": utc_now_iso()}, merge=True)
    except Exception:
        pass

    return redirect(url_for("admin_users", token=token, message=f"Plan updated for user {user_id} -> {desired}"))


@app.route("/admin/users/<int:user_id>/limit", methods=["POST"])
def admin_user_set_limit(user_id: int):
    ok, resp = admin_required()
    if not ok:
        if (os.getenv("ADMIN_TOKEN") or "").strip():
            return resp, 403
        return resp, 404
    token = (request.args.get("token", "") or "").strip()

    raw = str(request.form.get("limit") or "").strip()
    if raw == "":
        try:
            set_user_daily_limit_override(int(user_id), None)
        except Exception:
            return redirect(url_for("admin_users", token=token, message="Failed to clear limit override."))
        return redirect(url_for("admin_users", token=token, message=f"Limit override cleared for user {user_id}"))

    try:
        n = int(raw)
    except ValueError:
        return redirect(url_for("admin_users", token=token, message="Limit must be an integer (or empty)."))
    if n < 0 or n > 100000:
        return redirect(url_for("admin_users", token=token, message="Limit must be between 0 and 100000."))

    try:
        set_user_daily_limit_override(int(user_id), int(n))
    except Exception:
        return redirect(url_for("admin_users", token=token, message="Failed to update limit override."))

    return redirect(url_for("admin_users", token=token, message=f"Limit override set for user {user_id} -> {n}"))


@app.route("/admin/users/<int:user_id>/reset-usage", methods=["POST"])
def admin_user_reset_usage(user_id: int):
    ok, resp = admin_required()
    if not ok:
        if (os.getenv("ADMIN_TOKEN") or "").strip():
            return resp, 403
        return resp, 404
    token = (request.args.get("token", "") or "").strip()
    scope = str(request.form.get("scope") or "today").strip().lower()
    if scope not in {"today", "all"}:
        scope = "today"

    conn = get_conn()
    try:
        cur = conn.cursor()
        if scope == "all":
            cur.execute("DELETE FROM daily_message_usage WHERE user_id = ?", (int(user_id),))
        else:
            tz = get_user_tz(int(user_id)) or normalize_tz_name((os.getenv("USAGE_TZ") or "").strip()) or "UTC"
            day = usage_day_iso(tz)
            cur.execute("DELETE FROM daily_message_usage WHERE user_id = ? AND day = ?", (int(user_id), day))
        conn.commit()
    finally:
        conn.close()

    if scope != "all":
        try:
            firebase_uid = _firebase_uid_for_user_id(int(user_id))
            if firebase_uid and _billing_firestore_client() is not None:
                tz = get_user_tz(int(user_id)) or normalize_tz_name((os.getenv("USAGE_TZ") or "").strip()) or "UTC"
                day = usage_day_iso(tz)
                doc_ref = _billing_firestore_doc(firebase_uid)
                if doc_ref is not None:
                    doc_ref.set({"uid": firebase_uid, "last_reset_date": day, "messages_used_today": 0, "updated_at_utc": utc_now_iso()}, merge=True)
        except Exception:
            pass

    return redirect(url_for("admin_users", token=token, message=f"Usage reset ({scope}) for user {user_id}"))


@app.route("/admin/users/<int:user_id>/force-revoke-pro", methods=["POST"])
def admin_user_force_revoke_pro(user_id: int):
    ok, resp = admin_required()
    if not ok:
        if (os.getenv("ADMIN_TOKEN") or "").strip():
            return resp, 403
        return resp, 404
    token = (request.args.get("token", "") or "").strip()

    try:
        _set_user_plan(int(user_id), FREE_PLAN)
    except Exception:
        return redirect(url_for("admin_users", token=token, message="Failed to set user plan to free."))

    try:
        _billing_clear_subscription_local(int(user_id))
    except Exception:
        pass

    try:
        firebase_uid = _firebase_uid_for_user_id(int(user_id))
        if firebase_uid:
            _billing_force_revoke_firestore_entitlements(firebase_uid)
    except Exception:
        pass

    return redirect(url_for("admin_users", token=token, message=f"Pro entitlements revoked for user {user_id}"))


@app.route("/admin/feedback/delete", methods=["POST"])
def admin_feedback_delete():
    ok, resp = admin_required()
    if not ok:
        if (os.getenv("ADMIN_TOKEN") or "").strip():
            return resp, 403
        return resp, 404

    token = (request.args.get("token", "") or "").strip()
    limit_raw = (request.form.get("limit", "") or request.args.get("limit", "") or "").strip()
    try:
        limit = int(limit_raw) if limit_raw else 200
    except ValueError:
        limit = 200
    limit = max(1, min(2000, limit))

    ids_raw = (request.form.get("ids", "") or "").strip()
    fid_raw = (request.form.get("id", "") or "").strip()
    ids: list[int] = []
    if ids_raw:
        for part in ids_raw.split(","):
            p = (part or "").strip()
            if not p:
                continue
            try:
                ids.append(int(p))
            except ValueError:
                return redirect(url_for("admin_feedback", limit=limit, token=token, message="Invalid feedback id list."))
    elif fid_raw:
        try:
            ids = [int(fid_raw)]
        except ValueError:
            return redirect(url_for("admin_feedback", limit=limit, token=token, message="Invalid feedback id."))
    else:
        return redirect(url_for("admin_feedback", limit=limit, token=token, message="No feedback selected."))

    # Deduplicate and keep only positive IDs.
    ids = sorted({i for i in ids if i > 0})
    if not ids:
        return redirect(url_for("admin_feedback", limit=limit, token=token, message="No feedback selected."))

    conn = get_conn()
    try:
        cur = conn.cursor()
        if len(ids) == 1:
            cur.execute("DELETE FROM feedback WHERE id = ?", (ids[0],))
        else:
            placeholders = ",".join(["?"] * len(ids))
            cur.execute(f"DELETE FROM feedback WHERE id IN ({placeholders})", tuple(ids))
        conn.commit()
        deleted = int(getattr(cur, "rowcount", 0) or 0)
        if deleted > 0:
            if len(ids) == 1:
                msg = f"Deleted feedback #{ids[0]}."
            else:
                msg = f"Deleted {deleted} feedback item(s)."
        else:
            if len(ids) == 1:
                msg = f"Feedback #{ids[0]} was not found."
            else:
                msg = "No feedback items were deleted."
    except sqlite3.OperationalError as exc:
        msg = f"DB error: {exc}"
    finally:
        conn.close()

    return redirect(url_for("admin_feedback", limit=limit, token=token, message=msg))


@app.route("/admin/feedback.csv", methods=["GET"])
def admin_feedback_csv():
    ok, resp = admin_required()
    if not ok:
        if (os.getenv("ADMIN_TOKEN") or "").strip():
            return resp, 403
        return resp, 404

    limit_raw = (request.args.get("limit", "") or "").strip()
    try:
        limit = int(limit_raw) if limit_raw else 2000
    except ValueError:
        limit = 2000
    limit = max(1, min(10000, limit))

    import csv
    import io

    conn = get_conn()
    try:
        cur = conn.cursor()
        cur.execute(
            """
            SELECT id, created_at_utc, user_id, COALESCE(username, ''), COALESCE(contact_email, ''), COALESCE(category, ''),
                   COALESCE(rating, ''), COALESCE(message, ''), COALESCE(page, ''), COALESCE(tz, ''),
                   COALESCE(user_agent, ''), COALESCE(ip, '')
            FROM feedback
            ORDER BY id DESC
            LIMIT ?
            """,
            (limit,),
        )
        output = io.StringIO()
        w = csv.writer(output)
        w.writerow(
            [
                "id",
                "created_at_utc",
                "user_id",
                "username",
                "contact_email",
                "category",
                "rating",
                "message",
                "page",
                "tz",
                "user_agent",
                "ip",
            ]
        )
        for r in cur.fetchall():
            w.writerow(list(r))
        csv_text = output.getvalue()
    finally:
        conn.close()

    return Response(
        csv_text,
        mimetype="text/csv",
        headers={"Content-Disposition": 'attachment; filename="feedback.csv"'},
    )


if socketio is not None:
    @socketio.on("connect")
    def sio_connect():
        if "user_id" not in session:
            return False
        user_id = int(session["user_id"])
        join_room(f"user:{user_id}")

    @socketio.on("conversation:join")
    def sio_conversation_join(data):
        if "user_id" not in session:
            return
        user_id = int(session["user_id"])
        conversation_id = int((data or {}).get("conversation_id") or 0)
        if not conversation_id or not user_is_member(conversation_id, user_id):
            return
        join_room(f"conv:{conversation_id}")

    @socketio.on("message:send")
    def sio_message_send(data):
        if "user_id" not in session:
            return
        user_id = int(session["user_id"])
        conversation_id = int((data or {}).get("conversation_id") or 0)
        body = str((data or {}).get("body") or "").strip()
        if not conversation_id or not body:
            return
        if not user_is_member(conversation_id, user_id):
            return

        conn = get_conn()
        try:
            cur = conn.cursor()
            created_at = utc_now_iso()
            cur.execute(
                """
                INSERT INTO chat_messages (conversation_id, sender_id, type, body, media_path, created_at_utc)
                VALUES (?, ?, 'text', ?, '', ?)
                """,
                (conversation_id, user_id, body, created_at),
            )
            msg_id = int(cur.lastrowid)
            conn.commit()
        finally:
            conn.close()

        msg = {
            "id": msg_id,
            "conversation_id": conversation_id,
            "sender_id": user_id,
            "type": "text",
            "body": body,
            "media_path": "",
            "created_at_utc": created_at,
        }

        emit("message:new", msg, to=f"conv:{conversation_id}")

        # Update conversation list / unread counters for each member.
        for mid in conversation_member_ids(conversation_id):
            brief = conversation_brief_for_user(conversation_id, mid)
            if brief:
                emit("conversation:update", brief, to=f"user:{mid}")

    @socketio.on("conversation:read")
    def sio_conversation_read(data):
        if "user_id" not in session:
            return
        user_id = int(session["user_id"])
        conversation_id = int((data or {}).get("conversation_id") or 0)
        if not conversation_id:
            return
        mark_conversation_read(conversation_id, user_id)
        # Notify peers to refresh read state.
        for mid in conversation_member_ids(conversation_id):
            brief = conversation_brief_for_user(conversation_id, mid)
            if brief:
                emit("conversation:update", brief, to=f"user:{mid}")

    # Calls (WebRTC signaling over Socket.IO). This is P2P, no TURN server included.
    @socketio.on("call:invite")
    def sio_call_invite(data):
        if "user_id" not in session:
            return
        user_id = int(session["user_id"])
        to_user_id = int((data or {}).get("to_user_id") or 0)
        call_id = str((data or {}).get("call_id") or "").strip()
        video = bool((data or {}).get("video") or False)
        if not to_user_id or not call_id or to_user_id == user_id:
            return
        me = get_user_brief(user_id) or {"id": user_id, "username": session.get("username", "User"), "photo_path": ""}
        emit(
            "call:incoming",
            {"call_id": call_id, "from": me, "video": video},
            to=f"user:{to_user_id}",
        )

    @socketio.on("call:accept")
    def sio_call_accept(data):
        if "user_id" not in session:
            return
        user_id = int(session["user_id"])
        call_id = str((data or {}).get("call_id") or "").strip()
        to_user_id = int((data or {}).get("to_user_id") or 0)
        if not call_id or not to_user_id:
            return
        me = get_user_brief(user_id) or {"id": user_id, "username": session.get("username", "User"), "photo_path": ""}
        emit("call:accepted", {"call_id": call_id, "by": me}, to=f"user:{to_user_id}")

    @socketio.on("call:decline")
    def sio_call_decline(data):
        if "user_id" not in session:
            return
        user_id = int(session["user_id"])
        call_id = str((data or {}).get("call_id") or "").strip()
        to_user_id = int((data or {}).get("to_user_id") or 0)
        if not call_id or not to_user_id:
            return
        emit("call:declined", {"call_id": call_id, "by_user_id": user_id}, to=f"user:{to_user_id}")

    @socketio.on("call:end")
    def sio_call_end(data):
        if "user_id" not in session:
            return
        user_id = int(session["user_id"])
        call_id = str((data or {}).get("call_id") or "").strip()
        to_user_id = int((data or {}).get("to_user_id") or 0)
        if not call_id or not to_user_id:
            return
        emit("call:ended", {"call_id": call_id, "by_user_id": user_id}, to=f"user:{to_user_id}")

    @socketio.on("webrtc:offer")
    def sio_webrtc_offer(data):
        if "user_id" not in session:
            return
        to_user_id = int((data or {}).get("to_user_id") or 0)
        if not to_user_id:
            return
        emit("webrtc:offer", data, to=f"user:{to_user_id}")

    @socketio.on("webrtc:answer")
    def sio_webrtc_answer(data):
        if "user_id" not in session:
            return
        to_user_id = int((data or {}).get("to_user_id") or 0)
        if not to_user_id:
            return
        emit("webrtc:answer", data, to=f"user:{to_user_id}")

    @socketio.on("webrtc:ice")
    def sio_webrtc_ice(data):
        if "user_id" not in session:
            return
        to_user_id = int((data or {}).get("to_user_id") or 0)
        if not to_user_id:
            return
        emit("webrtc:ice", data, to=f"user:{to_user_id}")


if __name__ == "__main__":
    host = os.getenv("HOST", "0.0.0.0")
    port = int(os.getenv("PORT", "5000"))
    debug = os.getenv("FLASK_DEBUG", "0") == "1"
    if socketio is not None:
        # Flask-SocketIO blocks running with Werkzeug in "production" unless explicitly allowed.
        socketio.run(app, host=host, port=port, debug=debug, allow_unsafe_werkzeug=True)
    else:
        app.run(host=host, port=port, debug=debug)
